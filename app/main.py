import asyncio
import re
import secrets
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
import io
import openpyxl

from app.config import settings
from app.database import db, init_db, utcnow_iso
from app.services.reminders import due_reminders, get_registration, get_reminder_settings, has_sent_reminder, mark_sent, schedule_event_reminder, schedule_reminder
from app.services.sendpulse import upsert_contact as sendpulse_upsert
from app.services.telegram_api import (
    answer_callback,
    back_button,
    delete_message,
    delete_webhook,
    edit_message_text,
    get_updates,
    inline_keyboard,
    inline_url_button,
    main_inline_menu,
    main_menu_keyboard,
    phone_keyboard,
    remove_keyboard,
    send_message,
    send_document,
    send_photo,
    send_video,
)
from app.services.wayforpay import check_payment_status, create_invoice, success_ack, verify_callback
from pathlib import Path
from urllib.parse import quote, urlencode

UPLOADS_DIR = Path(__file__).resolve().parent.parent / "uploads"
UPLOADS_DIR.mkdir(exist_ok=True)


def _parse_datetime_value(value: str | datetime | None) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    if not s:
        return None
    s = s.replace(" ", "T")
    if s.endswith("Z"):
        s = s[:-1] + "+00:00"
    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None


def _format_datetime(value: str | datetime | None, tz_name: str = "Europe/Kyiv") -> str:
    """Перетворити datetime у читабельний формат dd.mm.yyyy HH:MM (час Києва для aware datetime)."""
    dt = _parse_datetime_value(value)
    if not dt:
        return str(value or "")
    if dt.tzinfo is not None:
        try:
            dt = dt.astimezone(ZoneInfo(tz_name))
        except Exception:
            pass
    return dt.strftime("%d.%m.%Y %H:%M")


def _format_datetime_input(value: str | datetime | None, tz_name: str = "Europe/Kyiv") -> str:
    """Формат для input[type=datetime-local]: YYYY-MM-DDTHH:MM."""
    dt = _parse_datetime_value(value)
    if not dt:
        return ""
    if dt.tzinfo is not None:
        try:
            dt = dt.astimezone(ZoneInfo(tz_name))
        except Exception:
            pass
    return dt.strftime("%Y-%m-%dT%H:%M")


async def _save_uploaded_media(upload: UploadFile | None, prefix: str) -> tuple[str, str]:
    """Save uploaded media and return (public_url, detected_content_type)."""
    if not upload or not upload.filename:
        return "", ""
    data = await upload.read()
    if not data or len(data) > 50 * 1024 * 1024:
        return "", ""
    src_mime = (upload.content_type or "").lower()
    ext = Path(upload.filename or "file").suffix or ".bin"
    if src_mime.startswith("image/"):
        ctype = "photo"
    elif src_mime.startswith("video/"):
        ctype = "video"
    else:
        ctype = "document"
    fname = f"{prefix}_{int(datetime.now(timezone.utc).timestamp())}_{secrets.token_hex(4)}{ext}"
    (UPLOADS_DIR / fname).write_bytes(data)
    media_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
    return media_url, ctype


def _get_app_config(key: str, default: str = "") -> str:
    """Отримати значення з app_config (fallback на default)."""
    with db() as conn:
        row = conn.execute("SELECT value FROM app_config WHERE key = ?", (key,)).fetchone()
    return (row["value"] or "").strip() if row else default


def _set_app_config(key: str, value: str) -> None:
    with db() as conn:
        conn.execute("INSERT OR REPLACE INTO app_config (key, value) VALUES (?, ?)", (key, (value or "").strip()))


def _parse_tg_restrict_ids(raw: str) -> set[int] | None:
    """Рядки / коми з числовими Telegram user id; порожньо — без обмеження."""
    if not (raw or "").strip():
        return None
    out: set[int] = set()
    for part in raw.replace(",", "\n").split("\n"):
        part = part.strip()
        if not part:
            continue
        try:
            out.add(int(part))
        except ValueError:
            continue
    return out or None


def audit_log(kind: str, user_id: int | None = None, event_id: int | None = None, detail: str = "") -> None:
    try:
        with db() as conn:
            conn.execute(
                "INSERT INTO audit_log (created_at, kind, user_id, event_id, detail) VALUES (?,?,?,?,?)",
                (utcnow_iso(), kind, user_id, event_id, (detail or "")[:2000]),
            )
    except Exception:
        pass


def _cfg_int(key: str, default: int = 0) -> int:
    try:
        v = (_get_app_config(key, str(default)) or "").strip()
        return int(v) if v != "" else default
    except ValueError:
        return default


def _timeline_send_allowed(kind: str) -> bool:
    key = {"rsvp_40": "timeline_rsvp_40_enabled", "24h": "timeline_24h_enabled", "3h": "timeline_3h_enabled"}.get(kind, "timeline_3h_enabled")
    return _cfg_int(key, 1) == 1


def _is_quiet_hours_now() -> bool:
    if _cfg_int("quiet_hours_enabled", 0) != 1:
        return False
    try:
        tzname = (_get_app_config("quiet_hours_tz", "Europe/Kyiv") or "").strip() or "Europe/Kyiv"
        tz = ZoneInfo(tzname)
    except Exception:
        return False
    h = datetime.now(tz).hour
    start = _cfg_int("quiet_hours_start", 22)
    end = _cfg_int("quiet_hours_end", 8)
    if start > end:
        return h >= start or h < end
    return start <= h < end


def _format_event_reminder_body(ev: dict, kind: str, title: str, dt_human: str, loc: str, zoom: str) -> str:
    col = {"rsvp_40": "msg_rsvp_40h", "24h": "msg_remind_24h", "3h": "msg_remind_3h"}[kind]
    raw = (ev.get(col) or "").strip()
    dkey = {"rsvp_40": "default_msg_rsvp_40h", "24h": "default_msg_remind_24h", "3h": "default_msg_remind_3h"}[kind]
    tpl = raw or (_get_app_config(dkey, "") or "").strip()
    if not tpl:
        tpl = {
            "rsvp_40": "Нагадування: подія «{title}» ({datetime}). Підтвердіть, чи будете присутні?",
            "24h": "Нагадування (за добу): «{title}»\n📅 {datetime}\n📍 {location}\n🔗 {zoom}",
            "3h": "Через кілька годин: «{title}» ({datetime}).",
        }[kind]
    return tpl.replace("{title}", title).replace("{datetime}", dt_human).replace("{location}", loc).replace("{zoom}", zoom)


async def _process_waitlist_for_event(event_id: int) -> None:
    """Один наступний у waitlist отримує повідомлення, якщо є вільне місце."""
    with db() as conn:
        ev = conn.execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
        if not ev:
            return
        evd = dict(ev)
        if not int(evd.get("waitlist_enabled") if evd.get("waitlist_enabled") is not None else 1):
            return
        lim = int(evd.get("seats_limit") or 0)
        if not lim:
            return
        seat_sql = _sql_registration_occupies_seat()
        occ = conn.execute(
            f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
            (event_id,),
        ).fetchone()["c"]
        if occ >= lim:
            return
        w = conn.execute(
            """SELECT w.id, w.user_id, u.tg_user_id FROM event_waitlist w
            JOIN users u ON u.id = w.user_id
            WHERE w.event_id = ? AND w.notified_at IS NULL
            ORDER BY w.created_at ASC LIMIT 1""",
            (event_id,),
        ).fetchone()
        if not w:
            return
        wid, wuid, tg = w["id"], w["user_id"], w["tg_user_id"]
    await send_message(
        int(tg),
        "Звільнилося місце на події! Можете зареєструватись, поки воно вільне.",
        reply_markup=inline_keyboard([[{"text": "Відкрити подію", "callback_data": f"evd_{event_id}"}]]),
    )
    with db() as conn:
        conn.execute("UPDATE event_waitlist SET notified_at = ? WHERE id = ?", (utcnow_iso(), wid))
    audit_log("waitlist_notified", user_id=wuid, event_id=event_id)


def registration_tz_label(r: dict, event_d: dict) -> str:
    """Мітка статусу за ТЗ §2.2 (на основі поточних полів БД)."""
    ps = (r.get("payment_status") or "").strip()
    st = (r.get("status") or "").strip()
    rsvp = r.get("rsvp_will_attend")
    attended = int(r.get("attended") or 0)
    fs = int(r.get("feedback_sent") or 0)
    ns = int(r.get("no_show") or 0)
    archived = (event_d.get("status") or "").strip() == "archived"
    if fs:
        return "feedback_done"
    if attended:
        return "checked_in"
    if (archived or ns) and attended == 0 and ps in ("paid", "free") and st == "registered":
        return "no_show"
    if rsvp == 1 and ps in ("paid", "free") and st == "registered":
        return "confirmed_attendance"
    if rsvp == 0 and st == "registered":
        return "declined_attendance"
    if ps in ("paid", "free") and st == "registered":
        return "paid"
    if st == "pending_payment" or ps not in ("paid", "free"):
        return "pending_payment"
    if st == "registered":
        return "registered"
    return st or "—"


def _gcal_url(ev: dict) -> str:
    """Побудувати URL для додавання події в Google Календар."""
    try:
        start_dt = datetime.fromisoformat((ev.get("starts_at") or "").replace("Z", "+00:00")[:19])
        if start_dt.tzinfo is None:
            start_dt = start_dt.replace(tzinfo=timezone.utc)
        start_utc = start_dt.astimezone(timezone.utc)
        end_utc = start_utc + timedelta(hours=2)
        dates = f"{start_utc.strftime('%Y%m%dT%H%M%SZ')}/{end_utc.strftime('%Y%m%dT%H%M%SZ')}"
    except Exception:
        dates = f"{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}/{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}"
    title = quote(str(ev.get("title") or "Подія")[:100])
    details = quote(str(ev.get("description") or ev.get("location") or "")[:200])
    location = quote(str(ev.get("location") or "")[:100])
    return f"https://calendar.google.com/calendar/render?action=TEMPLATE&text={title}&dates={dates}&details={details}&location={location}"


def _event_payment_url_from_template(
    template_url: str,
    *,
    order_ref: str,
    amount: int,
    event_id: int,
    user_id: int,
) -> str:
    url = (template_url or "").strip()
    if not url.startswith("http"):
        return ""
    resolved = (
        url.replace("{order_ref}", str(order_ref))
        .replace("{amount}", str(amount))
        .replace("{event_id}", str(event_id))
        .replace("{user_id}", str(user_id))
    )
    if "{order_ref}" not in url:
        sep = "&" if "?" in resolved else "?"
        resolved = f"{resolved}{sep}order_ref={quote(str(order_ref))}"
    return resolved


def _strike_text(s: str) -> str:
    return "".join((ch + "\u0336") for ch in str(s))


def _event_prices_for_role(event: dict, role: str) -> tuple[int, int | None]:
    role = (role or "guest").strip()
    if role == "member":
        base = int(event.get("member_price_uah") or 0)
        disc_raw = event.get("member_discount_price_uah")
    elif role == "lead":
        lead_base = event.get("lead_price_uah")
        base = int(lead_base if lead_base is not None else (event.get("guest_price_uah") or 0))
        disc_raw = event.get("lead_discount_price_uah")
    else:
        base = int(event.get("guest_price_uah") or 0)
        disc_raw = event.get("guest_discount_price_uah")
    discount: int | None = None
    if disc_raw is not None and str(disc_raw).strip() != "":
        try:
            discount = int(disc_raw)
        except Exception:
            discount = None
    return base, discount


async def _build_event_pay_url(user: dict, event: dict, reg: dict) -> tuple[str, str, int]:
    """Повертає URL оплати події, order_ref та amount."""
    role = (user.get("role") or "guest").strip()
    base_price, discount_price = _event_prices_for_role(event, role)
    amount = discount_price if discount_price is not None else base_price
    event_id = int(event.get("id") or 0)
    user_id = int(user.get("id") or 0)
    order_ref = (reg.get("wayforpay_order_reference") or "").strip() or f"ev{event_id}_u{user_id}_{int(datetime.now().timestamp())}"
    if not (reg.get("wayforpay_order_reference") or "").strip():
        with db() as conn:
            conn.execute(
                "UPDATE registrations SET wayforpay_order_reference = ?, amount_uah = ?, updated_at = ? WHERE user_id = ? AND event_id = ?",
                (order_ref, amount, utcnow_iso(), user_id, event_id),
            )
    if role == "member":
        tpl = event.get("wayforpay_member_url")
    elif role == "lead":
        tpl = event.get("wayforpay_lead_url")
    else:
        tpl = event.get("wayforpay_guest_url")
    pay_url = _event_payment_url_from_template(
        tpl or "",
        order_ref=order_ref,
        amount=int(amount or 0),
        event_id=event_id,
        user_id=user_id,
    )
    if pay_url:
        return pay_url, order_ref, amount
    invoice_url = await create_invoice(
        order_ref,
        float(amount),
        f"Квиток: {event.get('title','')}",
        client_email=user.get("email") or "",
        client_phone=user.get("phone") or "",
        client_name=user.get("full_name") or "",
    )
    return (invoice_url or ""), order_ref, amount


def _membership_payment_url_from_template(
    template_url: str,
    *,
    order_ref: str,
    amount: int,
    user_id: int,
    period: str,
) -> str:
    url = (template_url or "").strip()
    if not url.startswith("http"):
        return ""
    resolved = (
        url.replace("{order_ref}", str(order_ref))
        .replace("{amount}", str(amount))
        .replace("{user_id}", str(user_id))
        .replace("{period}", str(period))
    )
    if "{order_ref}" not in url:
        sep = "&" if "?" in resolved else "?"
        resolved = f"{resolved}{sep}order_ref={quote(str(order_ref))}"
    return resolved


def _resolve_templates_dir() -> Path:
    """Каталог `templates` у корені проєкту (ігнорує випадкові пробіли в назві папки)."""
    root = Path(__file__).resolve().parent.parent
    direct = root / "templates"
    if direct.is_dir():
        return direct
    for child in root.iterdir():
        if child.is_dir() and child.name.strip() == "templates":
            return child
    raise RuntimeError(
        f"Не знайдено папку шаблонів «templates» у {root}. "
        "Має бути каталог templates з HTML-шаблонами адмінки."
    )


app = FastAPI(title="Community Bot")
templates = Jinja2Templates(directory=str(_resolve_templates_dir()))
templates.env.filters["dt_kyiv"] = lambda v: _format_datetime(v, "Europe/Kyiv")
templates.env.filters["dt_input_kyiv"] = lambda v: _format_datetime_input(v, "Europe/Kyiv")
try:
    from fastapi.staticfiles import StaticFiles
    app.mount("/uploads", StaticFiles(directory=str(UPLOADS_DIR)), name="uploads")
except Exception:
    pass
signer = URLSafeTimedSerializer(settings.app_secret_key)

ADMIN_TOKEN_MAX_AGE = 86400 * 7  # 7 днів


def _decode_signed_token(token: str) -> dict | None:
    if not (token or "").strip():
        return None
    try:
        return signer.loads(token.strip(), max_age=ADMIN_TOKEN_MAX_AGE)
    except (BadSignature, SignatureExpired):
        return None


def _verify_admin_token(token: str) -> None:
    d = _decode_signed_token(token)
    if not d or "tg_user_id" not in d:
        raise HTTPException(status_code=403, detail="Invalid token")


def _verify_support_route_token(token: str, thread_id: int) -> None:
    d = _decode_signed_token(token)
    if not d:
        raise HTTPException(status_code=403, detail="Invalid token")
    if d.get("purpose") == "support_thread" and int(d.get("thread_id", -1)) == int(thread_id):
        return
    if "tg_user_id" in d:
        return
    raise HTTPException(status_code=403, detail="Invalid token")


def _make_support_thread_st_token(thread_id: int) -> str:
    # max_age застосовується на етапі loads(), а не dumps()
    return signer.dumps(
        {"purpose": "support_thread", "thread_id": int(thread_id), "nonce": secrets.token_hex(4)}
    )


def _admin_support_open_url_st(thread_id: int) -> str:
    st = _make_support_thread_st_token(thread_id)
    base = (settings.admin_public_url or settings.app_base_url).rstrip("/")
    return f"{base}/admin/support?{urlencode({'st': st, 'thread_id': str(thread_id)})}"


async def _notify_admin_support_group_new_ticket(*, thread_id: int, user_row: dict, query_text: str) -> None:
    chat_id = settings.admin_support_notify_chat_id
    if not chat_id:
        print("[BOT] ADMIN_SUPPORT_NOTIFY_CHAT_ID не задано в .env — сповіщення про підтримку в Telegram-групу вимкнено.")
        return
    thread_topic_id = settings.admin_support_notify_thread_id or None
    when = datetime.now(ZoneInfo("Europe/Kyiv")).strftime("%d.%m.%Y %H:%M")
    u = user_row or {}
    body = (
        "Нове звернення (підтримка)\n\n"
        f"ПІБ: {u.get('full_name') or '—'}\n"
        f"ТГ: {format_tg_display(u)}\n"
        f"Дата час: {when}\n\n"
        f"Запит:\n{(query_text or '').strip()[:800]}"
    )
    url = _admin_support_open_url_st(thread_id)
    try:
        await send_message(
            int(chat_id),
            body,
            reply_markup=inline_keyboard([[{"text": "Перейти до чату", "url": url}]]),
            message_thread_id=thread_topic_id,
        )
    except Exception as e:
        print(f"[BOT] support notify chat: {e}")

# Зберігання message_id останніх повідомлень бота в онбордингу для подальшого видалення
_onboarding_bot_msg_ids: dict[int, list[int]] = {}
# Повідомлення бота в анкеті «Стати учасником» — щоб видаляти при переході далі
_application_msg_ids: dict[int, list[int]] = {}
# Повідомлення бота в support — для видалення при «На головне меню»
_support_msg_ids: dict[int, list[int]] = {}
# Після входу в режим підтримки (меню/кнопка) — наступне повідомлення користувача має отримати шаблон «Дякуємо…»
_support_session_ack: dict[int, bool] = {}

SUPPORT_USER_ACK_TEXT = (
    "Дякуємо, за ваше звернення!\nАдмін напише вам найближчим часом в цьому чаті."
)
SUPPORT_USER_GOODBYE_TEXT = "Дякуємо за звернення!\nНа все добре 😊"


async def _delete_onboarding_msgs(chat_id: int) -> None:
    """Видалити збережені повідомлення бота з чату."""
    ids = _onboarding_bot_msg_ids.pop(chat_id, [])
    for mid in ids:
        try:
            await delete_message(chat_id, mid)
        except Exception:
            pass


async def _delete_application_msgs(chat_id: int) -> None:
    """Видалити промпти анкети «Стати учасником»."""
    ids = _application_msg_ids.pop(chat_id, [])
    for mid in ids:
        try:
            await delete_message(chat_id, mid)
        except Exception:
            pass


def get_or_create_user(tg_user_id: int, tg_username: str | None = None):
    with db() as conn:
        user = conn.execute("SELECT * FROM users WHERE tg_user_id = ?", (tg_user_id,)).fetchone()
        if user:
            if tg_username is not None:
                conn.execute("UPDATE users SET tg_username = ?, updated_at = ? WHERE id = ?", (tg_username or "", utcnow_iso(), user["id"]))
                user = conn.execute("SELECT * FROM users WHERE tg_user_id = ?", (tg_user_id,)).fetchone()
            return user
        now = utcnow_iso()
        conn.execute(
            "INSERT INTO users (tg_user_id, tg_username, role, created_at, updated_at) VALUES (?, ?, 'lead', ?, ?)",
            (tg_user_id, tg_username or "", now, now),
        )
        return conn.execute("SELECT * FROM users WHERE tg_user_id = ?", (tg_user_id,)).fetchone()


def set_stage(user_id: int, stage: str):
    with db() as conn:
        conn.execute("UPDATE users SET onboarding_stage = ?, updated_at = ? WHERE id = ?", (stage, utcnow_iso(), user_id))


def row_to_dict(row) -> dict:
    if row is None:
        return {}
    return {k: row[k] for k in row.keys()} if hasattr(row, "keys") else dict(row)


def format_tg_display(user_or_row: dict) -> str:
    """Повертає нік (@username) або fallback на tg_user_id."""
    u = user_or_row or {}
    username = (u.get("tg_username") or "").strip()
    if username:
        return f"@{username}" if not username.startswith("@") else username
    tid = u.get("tg_user_id")
    return f"tg:{tid}" if tid is not None else "—"


def ensure_open_support_thread(user_id: int) -> int:
    """Відкритий thread у БД з моменту входу в підтримку — інакше адмін не бачить чат, поки клієнт не надішле текст."""
    now = utcnow_iso()
    with db() as conn:
        row = conn.execute(
            "SELECT id FROM support_threads WHERE user_id = ? AND status = 'open' ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
        if row:
            conn.execute("UPDATE support_threads SET updated_at = ? WHERE id = ?", (now, row["id"]))
            return int(row["id"])
        conn.execute(
            "INSERT INTO support_threads (user_id, status, created_at, updated_at) VALUES (?, 'open', ?, ?)",
            (user_id, now, now),
        )
        new_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
    return int(new_id)


def parse_update(payload: dict):
    msg = payload.get("message") or payload.get("edited_message") or {}
    user = msg.get("from") or {}
    text = msg.get("text", "")
    contact = msg.get("contact")
    message_id = msg.get("message_id")
    return user, text, contact, message_id


def parse_start_deeplink(arg: str | None) -> tuple[int | None, str | None]:
    """Deep-link: `e12`, `event12`, опційно джерело `e12_facebook` (UTM/кампанія)."""
    if not arg or not str(arg).strip():
        return None, None
    s = str(arg).strip()
    m = re.match(r"^(?:e|event)(\d+)(?:_(.+))?$", s, re.IGNORECASE)
    if not m:
        return None, None
    src = (m.group(2) or "").strip()
    return int(m.group(1)), (src or None)


def _sql_registration_occupies_seat() -> str:
    return "status IN ('registered','pending_payment') AND (rsvp_will_attend IS NULL OR rsvp_will_attend != 0)"


def _set_user_deeplink(user_id: int, event_id: int | None, source: str | None) -> None:
    with db() as conn:
        conn.execute(
            "UPDATE users SET pending_deeplink_event_id = ?, deeplink_source = ?, updated_at = ? WHERE id = ?",
            (event_id, source or None, utcnow_iso(), user_id),
        )


async def render_event_detail(chat_id: int, user: dict, event_id: int, message_id: int | None = None) -> None:
    """Анонс події (як callback evd_) — для deep-link після онбордингу."""
    user_id = int(user["id"])
    role = (user.get("role") or "guest").strip()
    with db() as conn:
        ev = conn.execute("SELECT * FROM events WHERE id = ? AND status = 'upcoming'", (event_id,)).fetchone()
    if not ev:
        await send_message(chat_id, "Подію не знайдено або вона вже минула.")
        return
    ev = row_to_dict(ev)
    base_price, discount_price = _event_prices_for_role(ev, role)
    price = discount_price if discount_price is not None else base_price
    closed = " [закритий]" if ev.get("event_type") == "closed" else ""
    seat_sql = _sql_registration_occupies_seat()
    with db() as conn:
        reg_count = conn.execute(
            f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
            (event_id,),
        ).fetchone()["c"]
    seats_left = (ev.get("seats_limit") or 0) - reg_count if ev.get("seats_limit") else None
    loc = (ev.get("location") or "").strip()
    zoom = (ev.get("zoom_url") or "").strip()
    desc = (ev.get("description") or "").strip()[:500]
    msg = f"📌 {ev.get('title','')}{closed}\n\n📅 {_format_datetime(ev.get('starts_at') or '')}\n"
    if discount_price is not None:
        msg += f"💰 Ціна зі знижкою: {price} грн\n💸 Стара ціна: {_strike_text(str(base_price))} грн\n"
    else:
        msg += f"💰 Ціна: {price} грн\n"
    if seats_left is not None and ev.get("seats_limit"):
        msg += f"🪑 Залишилося місць: {seats_left}\n"
    if loc:
        msg += f"📍 {loc}\n"
    if zoom:
        msg += f"🔗 Zoom: {zoom}\n"
    if desc:
        msg += f"\n📝 {desc}\n"
    with db() as conn:
        reg = conn.execute("SELECT * FROM registrations WHERE user_id = ? AND event_id = ?", (user_id, event_id)).fetchone()
    btns: list = []
    if reg:
        regd = row_to_dict(reg)
        msg += "\n\n✅ Ви вже зареєстровані"
        if regd.get("payment_status") not in ("paid", "free"):
            pay_url, _, _ = await _build_event_pay_url(user, ev, regd)
            if pay_url:
                btns.append([{"text": "💳 Оплатити внесок", "url": pay_url}])
            else:
                btns.append([{"text": "💳 Оплатити внесок", "callback_data": f"evpay_{event_id}"}])
        if (
            regd.get("payment_status") in ("paid", "free")
            and regd.get("status") == "registered"
            and regd.get("rsvp_will_attend") is None
        ):
            btns.append(
                [
                    {"text": "✅ Буду", "callback_data": f"rsvp_yes_{event_id}"},
                    {"text": "❌ Не буду", "callback_data": f"rsvp_no_{event_id}"},
                ]
            )
        btns.append([{"text": "📅 Додати в Google Календар", "url": _gcal_url(ev)}])
    elif ev.get("event_type") == "closed" and role != "member":
        btns = [[{"text": "Стати учасником", "callback_data": "apply_member"}]]
    else:
        btns = [[{"text": "Зареєструватися", "callback_data": f"ev_{event_id}"}]]
    btns += back_button("events")
    markup = inline_keyboard(btns)
    if ev.get("banner_url"):
        media_url = str(ev["banner_url"]).lower()
        sent_ok = False
        try:
            if media_url.endswith((".mp4", ".mov", ".webm", ".m4v")):
                sent_ok = bool(await send_video(chat_id, ev["banner_url"], caption=msg, reply_markup=markup))
            else:
                sent_ok = bool(await send_photo(chat_id, ev["banner_url"], msg, reply_markup=markup))
        except Exception:
            sent_ok = False
        if sent_ok and message_id:
            await delete_message(chat_id, message_id)
        elif not sent_ok:
            await _edit_or_send_text(chat_id, message_id, msg, markup)
    else:
        await _edit_or_send_text(chat_id, message_id, msg, markup)


async def handle_start(tg_user_id: int, start_arg: str | None = None):
    chat_id = tg_user_id
    user = row_to_dict(get_or_create_user(tg_user_id))
    event_id, source = parse_start_deeplink(start_arg)
    if event_id is not None:
        with db() as conn:
            exists = conn.execute("SELECT 1 FROM events WHERE id = ? AND status = 'upcoming'", (event_id,)).fetchone()
        if exists:
            _set_user_deeplink(int(user["id"]), event_id, source)
        else:
            event_id = None
            _set_user_deeplink(int(user["id"]), None, None)
    has_profile = bool((user.get("full_name") or "").strip()) and bool((user.get("phone") or "").strip())
    if user.get("onboarding_stage") == "done" or has_profile:
        if user.get("onboarding_stage") != "done":
            # Якщо діалог з ботом видаляли/перезапускали, але профіль уже заповнений —
            # не запускаємо онбординг заново.
            set_stage(int(user["id"]), "done")
        role = user["role"]
        if event_id is not None:
            await render_event_detail(chat_id, user, event_id)
        else:
            msg = "🏠 Головне меню\n\n↓ Оберіть дію:"
            await send_message(chat_id, msg, reply_markup=inline_keyboard(main_inline_menu(role)))
        await send_message(chat_id, "\u200b", reply_markup=remove_keyboard())
        return
    await _delete_onboarding_msgs(chat_id)
    mid = await send_message(chat_id, "Вітаю! Введіть, будь ласка, ваше ПІБ.", return_message_id=True)
    if mid:
        _onboarding_bot_msg_ids.setdefault(chat_id, []).append(mid)
    set_stage(int(user["id"]), "waiting_full_name")
    schedule_reminder(int(user["id"]), "stage_1")


async def handle_text(tg_user_id: int, text: str, input_message_id: int | None = None):
    chat_id = tg_user_id
    user = row_to_dict(get_or_create_user(tg_user_id))
    if not user:
        await send_message(chat_id, "Помилка. Спробуйте /start")
        return
    stage = str(user.get("onboarding_stage") or "new")
    if input_message_id and stage != "waiting_support":
        try:
            await delete_message(chat_id, input_message_id)
        except Exception:
            pass
    user_id = int(user["id"])
    if stage == "waiting_full_name":
        with db() as conn:
            conn.execute("UPDATE users SET full_name = ?, onboarding_stage = ?, updated_at = ? WHERE id = ?", (text, "waiting_phone", utcnow_iso(), user_id))
        await _delete_onboarding_msgs(chat_id)
        mid = await send_message(chat_id, "Дякую! Натисніть кнопку, щоб поділитися номером.", reply_markup=phone_keyboard(), return_message_id=True)
        if mid:
            _onboarding_bot_msg_ids.setdefault(chat_id, []).append(mid)
        return
    if stage == "waiting_email":
        # Backward compatibility: users who were already in waiting_email
        # should continue onboarding without collecting email.
        with db() as conn:
            conn.execute("UPDATE users SET onboarding_stage = ?, updated_at = ? WHERE id = ?", ("waiting_phone", utcnow_iso(), user_id))
        await _delete_onboarding_msgs(chat_id)
        mid = await send_message(chat_id, "Натисніть кнопку, щоб поділитися номером.", reply_markup=phone_keyboard(), return_message_id=True)
        if mid:
            _onboarding_bot_msg_ids.setdefault(chat_id, []).append(mid)
        return

    with db() as conn:
        app_pending = conn.execute("SELECT * FROM member_application_pending WHERE user_id = ?", (user_id,)).fetchone()
    if app_pending:
        if app_pending["stage"] == "motivation":
            await _delete_application_msgs(chat_id)
            with db() as conn:
                conn.execute("INSERT INTO member_applications (user_id, motivation, created_at, status) VALUES (?, ?, ?, 'pending')", (user_id, text, utcnow_iso()))
                conn.execute("UPDATE member_application_pending SET stage = ? WHERE user_id = ?", ("experience", user_id))
            mid = await send_message(chat_id, "Дякую! Опишіть ваш досвід або інтереси:", return_message_id=True)
            if mid:
                _application_msg_ids[chat_id] = [mid]
            return
        if app_pending["stage"] == "experience":
            await _delete_application_msgs(chat_id)
            with db() as conn:
                app_row = conn.execute("SELECT motivation FROM member_applications WHERE user_id = ? AND status = 'pending'", (user_id,)).fetchone()
                conn.execute("UPDATE member_applications SET experience = ? WHERE user_id = ? AND status = 'pending'", (text, user_id))
                conn.execute("DELETE FROM member_application_pending WHERE user_id = ?", (user_id,))
            app_row_d = row_to_dict(app_row)
            motivation = (app_row_d.get("motivation") or "") if app_row_d else ""
            opys = (
                f"ПІБ: {user.get('full_name') or ''}\n"
                f"Телефон: {user.get('phone') or ''}\n"
                f"Email: {user.get('email') or ''}\n"
                f"Telegram: {format_tg_display(user)}\n\n"
                f"Мотивація: {motivation}\n\n"
                f"Досвід: {text}"
            ).strip()
            u = row_to_dict(user)
            asyncio.create_task(sendpulse_upsert({
                "external_id": str(u.get("id", "")), "full_name": u.get("full_name") or "", "phone": u.get("phone") or "",
                "email": u.get("email") or "", "role": "guest", "telegram_id": u.get("tg_user_id"),
                "tg_username": u.get("tg_username") or "", "payment_status": "member_application", "event_ids": "",
                "opys": opys,
            }))
            await _send_system_notice(chat_id, "Анкету отримано! Адміністратор зв'яжеться з вами для обговорення умов.")
            msg = "🏠 Головне меню\n\n↓ Оберіть дію:"
            await send_message(chat_id, msg, reply_markup=inline_keyboard(main_inline_menu("guest")))
            await send_message(chat_id, "\u200b", reply_markup=remove_keyboard())
            return

    if stage == "done" and text == "ℹ️ Про академію":
        body = (_get_app_config("academy_about_text", "") or "").strip() or "Текст ще не налаштований у адмін-панелі («Меню бота / інфо»)."
        await send_message(tg_user_id, body, reply_markup=inline_keyboard(back_button("main")))
        return
    if stage == "done" and text == "📄 Політика / оферта":
        url = (_get_app_config("policy_url", "") or "").strip()
        if url.startswith("http"):
            await send_message(tg_user_id, "Публічна оферта та політики:", reply_markup=inline_keyboard([[{"text": "📄 Відкрити документ", "url": url}]] + back_button("main")))
        else:
            await send_message(tg_user_id, "Посилання ще не вказане в адмін-панелі («Меню бота / інфо»).", reply_markup=inline_keyboard(back_button("main")))
        return
    if stage == "done" and text == "💬 Менеджер":
        raw = (_get_app_config("manager_contact", "") or "").strip()
        if raw.startswith("http"):
            await send_message(tg_user_id, "Зв'язок з менеджером:", reply_markup=inline_keyboard([[{"text": "💬 Відкрити", "url": raw}]] + back_button("main")))
        else:
            await send_message(tg_user_id, raw or "Контакт менеджера ще не налаштований («Меню бота / інфо»).", reply_markup=inline_keyboard(back_button("main")))
        return

    if text == "📅 Події" or text == "/events":
        now = datetime.now(timezone.utc)
        now_iso = now.isoformat()
        with db() as conn:
            near = conn.execute(
                "SELECT * FROM events WHERE status = 'upcoming' AND starts_at >= ? ORDER BY starts_at ASC LIMIT 10",
                (now_iso,),
            ).fetchall()
        msg = "📅 Найближчі події\n\n"
        btns = []
        if near:
            for e in near:
                msg += f"• {(e['title'] or '')[:45]} — {e['starts_at'][:16].replace('T', ' ')}\n"
                btns.append([{"text": f"📌 {(e['title'] or '')[:34]}", "callback_data": f"evd_{e['id']}"}])
        else:
            msg += "Найближчих подій поки немає.\n"
        msg += f"\n📆 Повний календар ({now.year}, {MONTH_NAMES[now.month - 1]})."
        btns.append([{"text": "📆 Календар по місяцях", "callback_data": "menu_events_cal"}])
        btns += back_button("main")
        await send_message(tg_user_id, msg, reply_markup=inline_keyboard(btns))
        return

    if text == "🏠 Мій кабінет":
        with db() as conn:
            regs = conn.execute(
                """SELECT r.*, e.title, e.starts_at, e.id as event_id FROM registrations r
                JOIN events e ON e.id = r.event_id WHERE r.user_id = ? AND r.status IN ('registered', 'pending_payment') AND e.status = 'upcoming'
                ORDER BY e.starts_at""",
                (user_id,),
            ).fetchall()
            paid_regs = conn.execute(
                """SELECT r.*, e.title, e.starts_at, e.id as event_id, e.status as event_status FROM registrations r
                JOIN events e ON e.id = r.event_id
                WHERE r.user_id = ? AND r.payment_status IN ('paid', 'free')
                ORDER BY e.starts_at DESC""",
                (user_id,),
            ).fetchall()
        role = (user.get("role") or "guest").strip()
        if role == "member":
            cab = "Кабінет Учасника"
        elif role == "lead":
            cab = "Кабінет Ліда"
        else:
            cab = "Кабінет Гостя"
        msg = f"🏠 {cab}\n\n"
        if role == "member":
            expires = (user.get("membership_expires_at") or "").strip()
            period = (user.get("membership_period") or "—").strip()
            amount = user.get("membership_amount_uah") or 0
            period_name = {"month": "місяць", "quarter": "квартал", "year": "рік"}.get(period, period)
            msg += f"📋 Статус: Учасник\n"
            msg += f"📅 Термін до: {expires[:10] if expires else '—'}\n"
            msg += f"💰 Наступний внесок: {amount} грн ({period_name})\n\n"
        else:
            role_name = "Гість" if role == "guest" else "Лід"
            msg += f"📋 Статус: {role_name}\n\n"
        btns = []
        if regs:
            msg += "📆 Майбутні заходи:\n\n"
            for r in regs:
                rd = row_to_dict(r) if hasattr(r, "keys") else r
                msg += f"  • {rd.get('title','')}\n    📅 {_format_datetime(rd.get('starts_at') or '')}\n\n"
        if role in ("guest", "lead"):
            msg += "✅ Оплачені події:\n\n"
            if paid_regs:
                for r in paid_regs:
                    rd = row_to_dict(r) if hasattr(r, "keys") else r
                    msg += f"  • {rd.get('title','')}\n    📅 {_format_datetime(rd.get('starts_at') or '')}\n\n"
            else:
                msg += "  Поки немає оплачених подій.\n\n"
        elif role == "member" and (user.get("membership_amount_uah") or 0) > 0:
            pass
        btns += back_button("main")
        await send_message(tg_user_id, msg, reply_markup=inline_keyboard(btns) if btns else None)
        return

    if text == "❓ Питання адміну":
        msg = "Напишіть ваше питання (воно буде переслано адміністратору). Пишіть далі — усі повідомлення потраплять до адміна."
        btns = [[{"text": "На головне меню", "callback_data": "support_to_main"}]]
        mid = await send_message(tg_user_id, msg, reply_markup=inline_keyboard(btns), return_message_id=True)
        if mid:
            _support_msg_ids[chat_id] = [mid]
        set_stage(user_id, "waiting_support")
        ensure_open_support_thread(user_id)
        _support_session_ack[user_id] = True
        return

    if stage == "waiting_support":
        now = utcnow_iso()
        u = row_to_dict(user)
        prev_user_msgs = 0
        thread_id = 0
        need_ack = _support_session_ack.pop(user_id, False)
        should_ack = False
        with db() as conn:
            thread = conn.execute(
                "SELECT id FROM support_threads WHERE user_id = ? AND status = 'open'",
                (user_id,),
            ).fetchone()
            if not thread:
                conn.execute(
                    "INSERT INTO support_threads (user_id, status, created_at, updated_at) VALUES (?, 'open', ?, ?)",
                    (user_id, now, now),
                )
                thread_id = int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])
            else:
                thread_id = int(thread["id"])
                conn.execute("UPDATE support_threads SET updated_at = ? WHERE id = ?", (now, thread_id))
            prev_user_msgs = conn.execute(
                "SELECT COUNT(*) as c FROM support_messages WHERE thread_id = ? AND from_admin = 0",
                (thread_id,),
            ).fetchone()["c"]
            conn.execute(
                "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 0, ?, 'text', ?)",
                (thread_id, text, now),
            )
            should_ack = bool(need_ack or (prev_user_msgs == 0))
            if should_ack:
                conn.execute(
                    "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, 'bot_outbound', ?)",
                    (thread_id, SUPPORT_USER_ACK_TEXT, now),
                )
        if should_ack:
            await send_message(chat_id, SUPPORT_USER_ACK_TEXT)
        try:
            await _notify_admin_support_group_new_ticket(thread_id=thread_id, user_row=u, query_text=text)
        except Exception as e:
            print(f"[BOT] support notify: {e}")
        tag_ids = []
        if settings.sendpulse_tag_contact_me:
            try:
                tag_ids = [int(settings.sendpulse_tag_contact_me.strip())]
            except ValueError:
                pass
        await sendpulse_upsert({
            "external_id": str(u.get("id", "")), "full_name": u.get("full_name") or "", "phone": u.get("phone") or "",
            "email": u.get("email") or "", "role": u.get("role") or "guest", "telegram_id": u.get("tg_user_id"),
            "tg_username": u.get("tg_username") or "", "payment_status": "", "event_ids": "",
            "opys": f"Питання адміну: {text}"[:500], "tag_ids": tag_ids,
        })
        return

    if text == "👤 Стати учасником" and user["role"] in ("guest", "lead"):
        with db() as conn:
            existing = conn.execute("SELECT 1 FROM member_applications WHERE user_id = ? AND status = 'pending'", (user_id,)).fetchone()
        if existing:
            await _send_system_notice(tg_user_id, "Ви вже подали анкету. Очікуйте зв'язку від адміна.")
            return
        with db() as conn:
            conn.execute("INSERT OR REPLACE INTO member_application_pending (user_id, stage) VALUES (?, 'motivation')", (user_id,))
        mid = await send_message(tg_user_id, "Розкажіть, чому ви хочете стати учасником спільноти:", return_message_id=True)
        if mid:
            _application_msg_ids[chat_id] = [mid]
        return

    if text == "👤 Моя картка" and user["role"] == "member":
        tg_link = f"https://t.me/{(user.get('tg_username') or '').strip()}" if (user.get("tg_username") or "").strip() else f"tg: {user.get('tg_user_id')}"
        with db() as conn:
            my_tags = conn.execute("SELECT t.id, t.name FROM member_tags t JOIN user_tags ut ON ut.tag_id = t.id WHERE ut.user_id = ?", (user_id,)).fetchall()
            all_tags = conn.execute("SELECT * FROM member_tags ORDER BY name").fetchall()
        my_ids = [t["id"] for t in my_tags]
        msg = f"👤 Картка Учасника\n\nПІБ: {user.get('full_name') or '—'}\n"
        msg += f"Telegram: {tg_link}\n\n"
        msg += "Теги: " + (", ".join(t["name"] for t in my_tags) if my_tags else "немає") + " (макс. 10)"
        available = [t for t in all_tags if t["id"] not in my_ids]
        if available and len(my_tags) < 10:
            btns = [[{"text": f"+ {t['name']}", "callback_data": f"addtag_{t['id']}"}] for t in available[:15]]
            await send_message(tg_user_id, msg + "\n\nДодати тег:", reply_markup=inline_keyboard(btns))
        else:
            await send_message(tg_user_id, msg + ("\nУсі теги додані." if len(my_tags) >= 10 and all_tags else "\nАдмін додасть теги в панелі." if not all_tags else ""))
        return

    if text == "👥 Картки учасників" and user["role"] == "member":
        with db() as conn:
            tags = conn.execute("SELECT * FROM member_tags").fetchall()
        if not tags:
            await send_message(tg_user_id, "Поки немає тегів. Адмін створить їх у панелі.")
            return
        btns = [[{"text": t["name"], "callback_data": f"tag_{t['id']}"}] for t in tags]
        await send_message(tg_user_id, "Оберіть тег для перегляду карток:", reply_markup=inline_keyboard(btns))
        return

    if text == "/admin":
        if tg_user_id not in settings.admin_ids:
            await send_message(tg_user_id, "У вас немає доступу до адмін-панелі.")
            return
        token = signer.dumps({"tg_user_id": tg_user_id, "nonce": secrets.token_hex(8)})
        base = (settings.admin_public_url or settings.app_base_url).rstrip("/")
        url = f"{base}/admin/entry/{token}"
        await send_message(
            tg_user_id,
            "Натисніть кнопку нижче, щоб відкрити адмін-панель:",
            reply_markup=inline_url_button("Відкрити адмін-панель", url),
        )
        return

    with db() as conn:
        pending = conn.execute("SELECT * FROM feedback_pending WHERE user_id = ?", (user["id"],)).fetchone()
    if pending:
        if pending["stage"] == "waiting_rating" and text.strip() in ("1", "2", "3", "4", "5"):
            rating = int(text.strip())
            with db() as conn:
                conn.execute(
                    "INSERT INTO feedback (user_id, event_id, rating, created_at) VALUES (?, ?, ?, ?)",
                    (user["id"], pending["event_id"], rating, utcnow_iso()),
                )
                conn.execute("UPDATE feedback_pending SET stage = ? WHERE user_id = ?", ("waiting_text", user["id"]))
            await send_message(tg_user_id, "Дякуємо! Залиште текстовий відгук (або напишіть /skip)")
            return
        if pending["stage"] == "waiting_text":
            review_text = "" if text.strip() == "/skip" else text
            with db() as conn:
                conn.execute(
                    "UPDATE feedback SET review_text = ? WHERE user_id = ? AND event_id = ?",
                    (review_text, user["id"], pending["event_id"]),
                )
                conn.execute("UPDATE registrations SET feedback_sent = 1 WHERE user_id = ? AND event_id = ?", (user["id"], pending["event_id"]))
                conn.execute("DELETE FROM feedback_pending WHERE user_id = ?", (user["id"],))
            audit_log("feedback_done", user_id=user["id"], event_id=pending["event_id"])
            await send_message(tg_user_id, "Дякуємо за відгук!")
            u = dict(user)
            await sendpulse_upsert({
                "external_id": str(u["id"]), "full_name": u.get("full_name") or "",
                "phone": u.get("phone") or "", "email": u.get("email") or "",
                "role": u.get("role") or "guest", "telegram_id": u.get("tg_user_id"),
                "tg_username": u.get("tg_username") or "", "payment_status": "",
                "event_ids": str(pending["event_id"]),
            })
            return

    if text.startswith("/register_event "):
        try:
            event_id = int(text.split(" ", 1)[1].strip())
        except Exception:
            await send_message(tg_user_id, "Формат: /register_event ID")
            return
        with db() as conn:
            event = conn.execute("SELECT * FROM events WHERE id = ? AND status = 'upcoming'", (event_id,)).fetchone()
            if not event:
                await send_message(tg_user_id, "Подію не знайдено.")
                return
            role = user["role"]
            if event["event_type"] == "closed" and role != "member":
                await send_message(tg_user_id, "Це закритий захід для учасників. Натисніть 'Стати учасником' у меню.")
                return
            lim = int(event["seats_limit"] or 0)
            if lim:
                seat_sql = _sql_registration_occupies_seat()
                occ = conn.execute(
                    f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
                    (event_id,),
                ).fetchone()["c"]
                mine = conn.execute(
                    f"SELECT 1 FROM registrations WHERE user_id = ? AND event_id = ? AND {seat_sql}",
                    (user["id"], event_id),
                ).fetchone()
                if occ >= lim and not mine:
                    evd = dict(event)
                    if int(evd.get("waitlist_enabled") if evd.get("waitlist_enabled") is not None else 1):
                        await send_message(
                            tg_user_id,
                            "На жаль, місць уже немає. Можете записатися в список очікування.",
                            reply_markup=inline_keyboard([[{"text": "Список очікування", "callback_data": f"wljoin_{event_id}"}]]),
                        )
                    else:
                        await send_message(tg_user_id, "На жаль, для цієї події вже немає вільних місць.")
                    return
            base_price, discount_price = _event_prices_for_role(event, role)
            amount = discount_price if discount_price is not None else base_price
            now = utcnow_iso()
            conn.execute(
                """
                INSERT OR REPLACE INTO registrations
                (user_id, event_id, status, amount_uah, payment_status, created_at, updated_at)
                VALUES (?, ?, 'pending_payment', ?, 'none', ?, ?)
                """,
                (user["id"], event_id, amount, now, now),
            )

        if amount == 0:
            with db() as conn:
                conn.execute(
                    "UPDATE registrations SET status = 'registered', payment_status = 'free', updated_at = ? WHERE user_id = ? AND event_id = ?",
                    (utcnow_iso(), user["id"], event_id),
                )
            await send_message(tg_user_id, "Ви успішно зареєстровані на подію.")
            schedule_event_reminder(user["id"], event_id)
            audit_log("register_free", user_id=user["id"], event_id=event_id)
            return

        order_ref = f"ev{event_id}_u{user['id']}_{int(datetime.now().timestamp())}"
        with db() as conn:
            conn.execute(
                "UPDATE registrations SET wayforpay_order_reference = ?, updated_at = ? WHERE user_id = ? AND event_id = ?",
                (order_ref, utcnow_iso(), user["id"], event_id),
            )
        u = dict(user)
        invoice_url = await create_invoice(
            order_ref, float(amount), f"Квиток: {event['title']}",
            client_email=u.get("email") or "",
            client_phone=u.get("phone") or "",
            client_name=u.get("full_name") or "",
        )
        if role == "member":
            tpl = event.get("wayforpay_member_url")
        elif role == "lead":
            tpl = event.get("wayforpay_lead_url")
        else:
            tpl = event.get("wayforpay_guest_url")
        pay_url = _event_payment_url_from_template(
            tpl or "",
            order_ref=order_ref,
            amount=int(amount or 0),
            event_id=event_id,
            user_id=int(user["id"]),
        ) or invoice_url
        if pay_url:
            pay_text = f"Натисніть кнопку нижче для оплати ({amount} грн):"
            if discount_price is not None:
                pay_text = (
                    f"Ціна зі знижкою: {amount} грн\n"
                    f"Стара ціна: {_strike_text(str(base_price))} грн\n\n"
                    "Натисніть кнопку нижче для оплати:"
                )
            await send_message(tg_user_id, pay_text, reply_markup=inline_url_button("💳 Перейти до оплати", pay_url))
        else:
            await send_message(tg_user_id, "На жаль, зараз неможливо створити платіжне посилання. Спробуйте пізніше або зверніться до адміна.")
        schedule_reminder(user["id"], "stage_4_unpaid_1h", event_id=event_id)
        schedule_reminder(user["id"], "stage_4_unpaid_24h", event_id=event_id)
        audit_log("register_pending_payment", user_id=user["id"], event_id=event_id)
        return

    if stage == "waiting_phone":
        await send_message(tg_user_id, "Натисніть кнопку внизу екрана, щоб поділитися номером телефону.")
        return
    if stage == "done":
        return
    await send_message(tg_user_id, "Скористайтесь меню бота або командою /start.")


async def _edit_or_send_text(chat_id: int, message_id: int | None, text: str, reply_markup: dict | None = None) -> None:
    """Спробувати editMessageText; якщо не вдалось (напр. повідомлення — фото) — надіслати нове."""
    if message_id:
        ok = await edit_message_text(chat_id, message_id, text, reply_markup=reply_markup)
        if ok:
            return
    await send_message(chat_id, text, reply_markup=reply_markup)


async def _send_system_notice(
    chat_id: int,
    text: str,
    reply_markup: dict | None = None,
    ttl_sec: int = 7200,
) -> None:
    """Системне повідомлення з авто-очисткою через ttl_sec секунд."""
    mid = await send_message(chat_id, text, reply_markup=reply_markup, return_message_id=True)
    if not mid or ttl_sec <= 0:
        return

    async def _cleanup() -> None:
        await asyncio.sleep(ttl_sec)
        try:
            await delete_message(chat_id, int(mid))
        except Exception:
            pass

    asyncio.create_task(_cleanup())


def _build_mycard_inline_view(user: dict) -> tuple[str, dict]:
    """Текст і inline-клавіатура для екрана «Моя картка» (те саме, що menu_mycard)."""
    tg_link = f"https://t.me/{(user.get('tg_username') or '').strip()}" if (user.get("tg_username") or "").strip() else f"tg: {user.get('tg_user_id')}"
    with db() as conn:
        my_tags = conn.execute(
            "SELECT t.id, t.name FROM member_tags t JOIN user_tags ut ON ut.tag_id = t.id WHERE ut.user_id = ?",
            (user["id"],),
        ).fetchall()
        all_tags = conn.execute("SELECT * FROM member_tags ORDER BY name").fetchall()
    msg = (
        f"👤 Моя картка\n\n📋 ПІБ: {user.get('full_name') or '—'}\n🔗 Telegram: {tg_link}\n\n🏷 Теги: "
        + (", ".join(t["name"] for t in my_tags) if my_tags else "немає")
        + " (макс. 10)"
    )
    available = [t for t in all_tags if t["id"] not in [x["id"] for x in my_tags]]
    btns: list = (
        [[{"text": f"+ {t['name']}", "callback_data": f"addtag_{t['id']}"}] for t in available[:15]]
        if available and len(my_tags) < 10
        else []
    )
    btns += back_button("main")
    return msg, inline_keyboard(btns)


MONTH_NAMES = ["Січень", "Лютий", "Березень", "Квітень", "Травень", "Червень", "Липень", "Серпень", "Вересень", "Жовтень", "Листопад", "Грудень"]


def _cal_buttons(y: int, m: int) -> list:
    """Кнопки кастомного календаря: рік, місяць, стрілки, події."""
    prev_y, prev_m = (y - 1, 12) if m == 1 else (y, m - 1)
    next_y, next_m = (y + 1, 1) if m == 12 else (y, m + 1)
    return [
        [{"text": "⬅️", "callback_data": f"cal_y_{prev_y}_{m}"}, {"text": str(y), "callback_data": "noop"}, {"text": "➡️", "callback_data": f"cal_y_{next_y}_{m}"}],
        [{"text": "⬅️", "callback_data": f"cal_m_{prev_y}_{prev_m:02d}"}, {"text": MONTH_NAMES[m - 1], "callback_data": "noop"}, {"text": "➡️", "callback_data": f"cal_m_{next_y}_{next_m:02d}"}],
        [{"text": f"📋 Події {MONTH_NAMES[m - 1].lower()} {y}", "callback_data": f"month_{y}_{m:02d}"}],
    ] + back_button("main")


async def handle_callback(cq_id: str, tg_user_id: int, chat_id: int, data: str, tg_username: str | None = None, message_id: int | None = None):
    await answer_callback(cq_id)
    user = row_to_dict(get_or_create_user(tg_user_id, tg_username=tg_username))
    if not user:
        return

    if data == "back_cards":
        with db() as conn:
            tags = conn.execute("SELECT * FROM member_tags").fetchall()
        if not tags:
            msg = "Поки немає тегів."
            btns = back_button("main")
        else:
            msg = "👥 Картки учасників\n\n↓ Оберіть тег:"
            btns = [[{"text": t["name"], "callback_data": f"tag_{t['id']}"}] for t in tags]
            btns += back_button("main")
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "back_main":
        role = user["role"]
        msg = "🏠 Головне меню\n\n↓ Оберіть дію:"
        btns = main_inline_menu(role)
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "refresh_member_menu":
        with db() as conn:
            u = conn.execute("SELECT * FROM users WHERE id = ?", (user["id"],)).fetchone()
        if u:
            u = dict(u)
            role_now = (u.get("role") or "guest").strip()
            msg = "🏠 Головне меню\n\n↓ Оберіть дію:"
            sent = await send_message(chat_id, msg, reply_markup=inline_keyboard(main_inline_menu(role_now)))
            if sent and message_id:
                await delete_message(chat_id, message_id)
        return

    if data == "menu_events":
        now = datetime.now(timezone.utc)
        now_iso = now.isoformat()
        with db() as conn:
            near = conn.execute(
                "SELECT * FROM events WHERE status = 'upcoming' AND starts_at >= ? ORDER BY starts_at ASC LIMIT 10",
                (now_iso,),
            ).fetchall()
        msg = "📅 Найближчі події\n\n"
        btns: list = []
        if near:
            for e in near:
                msg += f"• {(e['title'] or '')[:45]} — {e['starts_at'][:16].replace('T', ' ')}\n"
                btns.append([{"text": f"📌 {(e['title'] or '')[:34]}", "callback_data": f"evd_{e['id']}"}])
        else:
            msg += "Найближчих подій поки немає.\n"
        msg += f"\n📆 Повний календар ({now.year}, {MONTH_NAMES[now.month - 1]}) — натисніть нижче."
        btns.append([{"text": "📆 Календар по місяцях", "callback_data": "menu_events_cal"}])
        btns += back_button("main")
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "menu_events_cal":
        now = datetime.now(timezone.utc)
        msg = f"📅 Події\n\nРік: {now.year}\nМісяць: {MONTH_NAMES[now.month - 1]}\n\n↓ Оберіть місяць стрілками, потім натисніть для перегляду подій:"
        btns = _cal_buttons(now.year, now.month)
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "menu_cabinet":
        user_id = user["id"]
        with db() as conn:
            regs = conn.execute(
                """SELECT r.*, e.title, e.starts_at, e.id as event_id FROM registrations r
                JOIN events e ON e.id = r.event_id WHERE r.user_id = ? AND r.status IN ('registered', 'pending_payment') AND e.status = 'upcoming'
                ORDER BY e.starts_at""", (user_id,),
            ).fetchall()
            paid_regs = conn.execute(
                """SELECT r.*, e.title, e.starts_at, e.id as event_id, e.status as event_status FROM registrations r
                JOIN events e ON e.id = r.event_id
                WHERE r.user_id = ? AND r.payment_status IN ('paid', 'free')
                ORDER BY e.starts_at DESC""",
                (user_id,),
            ).fetchall()
        role = (user.get("role") or "guest").strip()
        if role == "member":
            cab = "Кабінет Учасника"
        elif role == "lead":
            cab = "Кабінет Ліда"
        else:
            cab = "Кабінет Гостя"
        msg = f"🏠 {cab}\n\n"
        if role == "member":
            expires = (user.get("membership_expires_at") or "").strip()
            period = (user.get("membership_period") or "—").strip()
            amount = user.get("membership_amount_uah") or 0
            period_name = {"month": "місяць", "quarter": "квартал", "year": "рік"}.get(period, period)
            msg += f"📋 Статус: Учасник\n📅 Термін до: {expires[:10] if expires else '—'}\n💰 Наступний внесок: {amount} грн ({period_name})\n\n"
        else:
            role_name = "Гість" if role == "guest" else "Лід"
            msg += f"📋 Статус: {role_name}\n\n"
        btns = []
        if regs:
            msg += "📆 Майбутні заходи:\n\n"
            for r in regs:
                rd = row_to_dict(r) if hasattr(r, "keys") else r
                msg += f"  • {rd.get('title','')}\n    📅 {_format_datetime(rd.get('starts_at') or '')}\n\n"
        if role in ("guest", "lead"):
            msg += "✅ Оплачені події:\n\n"
            if paid_regs:
                for r in paid_regs:
                    rd = row_to_dict(r) if hasattr(r, "keys") else r
                    msg += f"  • {rd.get('title','')}\n    📅 {_format_datetime(rd.get('starts_at') or '')}\n\n"
            else:
                msg += "  Поки немає оплачених подій.\n\n"
        elif role == "member" and (user.get("membership_amount_uah") or 0) > 0:
            pass
        btns += back_button("main")
        mk = inline_keyboard(btns) if btns else None
        await _edit_or_send_text(chat_id, message_id, msg, mk)
        return

    if data == "menu_mycard" and user["role"] == "member":
        msg, mk = _build_mycard_inline_view(user)
        await _edit_or_send_text(chat_id, message_id, msg, mk)
        return

    if data == "menu_cards" and user["role"] == "member":
        with db() as conn:
            tags = conn.execute("SELECT * FROM member_tags").fetchall()
        if not tags:
            msg = "Поки немає тегів."
            btns = back_button("main")
        else:
            msg = "👥 Картки учасників\n\n↓ Оберіть тег:"
            btns = [[{"text": t["name"], "callback_data": f"tag_{t['id']}"}] for t in tags]
            btns += back_button("main")
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "menu_apply" and user["role"] in ("guest", "lead"):
        with db() as conn:
            existing = conn.execute("SELECT 1 FROM member_applications WHERE user_id = ? AND status = 'pending'", (user["id"],)).fetchone()
        if existing:
            msg = "Ви вже подали анкету. Очікуйте зв'язку від адміна."
            btns = back_button("main")
        else:
            with db() as conn:
                conn.execute("INSERT OR REPLACE INTO member_application_pending (user_id, stage) VALUES (?, 'motivation')", (user["id"],))
            msg = "Розкажіть, чому ви хочете стати учасником спільноти:"
            btns = []  # без кнопки Назад — лише текст
        mk = inline_keyboard(btns) if btns else None
        if not existing:
            if message_id:
                ok = await edit_message_text(chat_id, message_id, msg, reply_markup=mk)
                if ok:
                    _application_msg_ids[chat_id] = [message_id]
                else:
                    nid = await send_message(chat_id, msg, reply_markup=mk, return_message_id=True)
                    if isinstance(nid, int):
                        _application_msg_ids[chat_id] = [nid]
            else:
                nid = await send_message(chat_id, msg, reply_markup=mk, return_message_id=True)
                if isinstance(nid, int):
                    _application_msg_ids[chat_id] = [nid]
        else:
            await _edit_or_send_text(chat_id, message_id, msg, mk)
        return

    if data == "menu_support":
        set_stage(user["id"], "waiting_support")
        ensure_open_support_thread(int(user["id"]))
        _support_session_ack[int(user["id"])] = True
        msg = "Напишіть ваше питання (воно буде переслано адміністратору). Пишіть далі — усі повідомлення потраплять до адміна."
        btns = [[{"text": "На головне меню", "callback_data": "support_to_main"}]]
        mk = inline_keyboard(btns)
        if message_id:
            ok = await edit_message_text(chat_id, message_id, msg, reply_markup=mk)
            if ok:
                _support_msg_ids[chat_id] = [message_id]
            else:
                mid = await send_message(chat_id, msg, reply_markup=mk, return_message_id=True)
                if isinstance(mid, int):
                    _support_msg_ids[chat_id] = [mid]
        else:
            mid = await send_message(chat_id, msg, reply_markup=mk, return_message_id=True)
            if isinstance(mid, int):
                _support_msg_ids[chat_id] = [mid]
        return

    if data == "menu_about":
        text = (_get_app_config("academy_about_text", "") or "").strip() or "Текст ще не налаштований. Увійдіть в адмін-панель → «Меню бота / інфо»."
        btns = back_button("main")
        await _edit_or_send_text(chat_id, message_id, text[:4090], inline_keyboard(btns))
        return

    if data == "menu_policy":
        url = (_get_app_config("policy_url", "") or "").strip()
        btns = back_button("main")
        if url.startswith("http"):
            msg = "Публічна оферта та політики — за посиланням нижче."
            await _edit_or_send_text(
                chat_id, message_id, msg, inline_keyboard([[{"text": "📄 Відкрити документ", "url": url}]] + back_button("main"))
            )
        else:
            msg = "Посилання на оферту ще не вказане в адмін-панелі («Меню бота / інфо»)."
            await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data == "menu_manager":
        raw = (_get_app_config("manager_contact", "") or "").strip()
        btns = back_button("main")
        if raw.startswith("http"):
            msg = "Зв'язок з менеджером:"
            row = [[{"text": "💬 Відкрити контакт", "url": raw}]]
            await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(row + back_button("main")))
        else:
            msg = raw or "Контакт менеджера ще не налаштований у адмін-панелі («Меню бота / інфо»)."
            await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data.startswith("rsvp_yes_"):
        eid = int(data[9:])
        with db() as conn:
            conn.execute(
                "UPDATE registrations SET rsvp_will_attend = 1, updated_at = ? WHERE user_id = ? AND event_id = ? AND status = 'registered'",
                (utcnow_iso(), user["id"], eid),
            )
        audit_log("rsvp_yes", user_id=user["id"], event_id=eid)
        await send_message(chat_id, "Підтверджено: будете на події ✅")
        await render_event_detail(chat_id, user, eid, message_id)
        return

    if data.startswith("rsvp_no_"):
        eid = int(data[8:])
        with db() as conn:
            conn.execute(
                "UPDATE registrations SET rsvp_will_attend = 0, updated_at = ? WHERE user_id = ? AND event_id = ? AND status = 'registered'",
                (utcnow_iso(), user["id"], eid),
            )
        await send_message(chat_id, "Дякуємо за відповідь. Місце звільнено для інших учасників.")
        audit_log("rsvp_no", user_id=user["id"], event_id=eid)
        await _process_waitlist_for_event(eid)
        await render_event_detail(chat_id, user, eid, message_id)
        return

    if data.startswith("wljoin_"):
        eid = int(data[7:])
        with db() as conn:
            exists = conn.execute("SELECT 1 FROM events WHERE id = ? AND status = 'upcoming'", (eid,)).fetchone()
            if not exists:
                await send_message(chat_id, "Подію не знайдено.")
                return
            conn.execute(
                "INSERT OR IGNORE INTO event_waitlist (user_id, event_id, created_at) VALUES (?, ?, ?)",
                (user["id"], eid, utcnow_iso()),
            )
        await send_message(chat_id, "Вас додано до списку очікування. Повідомимо, якщо звільниться місце.")
        audit_log("waitlist_join", user_id=user["id"], event_id=eid)
        return

    if data == "support_to_main":
        _support_session_ack.pop(int(user["id"]), None)
        set_stage(user["id"], "done")
        for mid in _support_msg_ids.pop(chat_id, []):
            try:
                await delete_message(chat_id, mid)
            except Exception:
                pass
        role = user.get("role") or "guest"
        msg = "🏠 Головне меню\n\n↓ Оберіть дію:"
        await send_message(chat_id, msg, reply_markup=inline_keyboard(main_inline_menu(role)))
        return

    if data == "back_events":
        now = datetime.now(timezone.utc)
        now_iso = now.isoformat()
        with db() as conn:
            near = conn.execute(
                "SELECT * FROM events WHERE status = 'upcoming' AND starts_at >= ? ORDER BY starts_at ASC LIMIT 10",
                (now_iso,),
            ).fetchall()
        msg = "📅 Найближчі події\n\n"
        btns: list = []
        if near:
            for e in near:
                msg += f"• {(e['title'] or '')[:45]} — {e['starts_at'][:16].replace('T', ' ')}\n"
                btns.append([{"text": f"📌 {(e['title'] or '')[:34]}", "callback_data": f"evd_{e['id']}"}])
        else:
            msg += "Найближчих подій поки немає.\n"
        msg += f"\n📆 Повний календар ({now.year}, {MONTH_NAMES[now.month - 1]}) — натисніть нижче."
        btns.append([{"text": "📆 Календар по місяцях", "callback_data": "menu_events_cal"}])
        btns += back_button("main")
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data.startswith("cal_y_") or data.startswith("cal_m_"):
        parts = data.split("_")
        y, m = int(parts[2]), int(parts[3])
        msg = f"📅 Події\n\nРік: {y}\nМісяць: {MONTH_NAMES[m - 1]}\n\n↓ Оберіть місяць стрілками, потім натисніть для перегляду подій:"
        btns = _cal_buttons(y, m)
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
        return

    if data.startswith("month_"):
        parts = data.split("_")
        y, m = int(parts[1]), int(parts[2])
        start = f"{y}-{m:02d}-01T00:00:00"
        end_m = m + 1 if m < 12 else 1
        end_y = y if m < 12 else y + 1
        end = f"{end_y}-{end_m:02d}-01T00:00:00"
        with db() as conn:
            rows = conn.execute(
                "SELECT * FROM events WHERE status = 'upcoming' AND starts_at >= ? AND starts_at < ? ORDER BY starts_at ASC LIMIT 20",
                (start, end),
            ).fetchall()
        if not rows:
            msg = f"📅 Події\n\nУ цьому місяці ({MONTH_NAMES[m - 1]} {y}) немає подій."
            btns = back_button("events")
            await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
            return
        role = user["role"]
        msg = f"📅 Події — {MONTH_NAMES[m - 1]} {y}\n\n"
        btns = []
        for e in rows[:15]:
            closed = " [закр.]" if e["event_type"] == "closed" else ""
            title = (e["title"] or "")[:40]
            msg += f"• {title}{closed} — {e['starts_at'][:10]}\n"
            btn_text = f"Опис події - {(e['title'] or '')[:35]}"
            btns.append([{"text": f"📌 {btn_text}{closed}", "callback_data": f"evd_{e['id']}"}])
        btns += back_button("events")
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
    elif data.startswith("evd_"):
        event_id = int(data[4:])
        await render_event_detail(chat_id, user, event_id, message_id)
    elif data.startswith("evpay_"):
        event_id = int(data[6:])
        with db() as conn:
            reg = conn.execute(
                "SELECT * FROM registrations WHERE user_id = ? AND event_id = ?",
                (user["id"], event_id),
            ).fetchone()
            event = conn.execute("SELECT * FROM events WHERE id = ? AND status = 'upcoming'", (event_id,)).fetchone()
        if not reg or not event:
            await send_message(chat_id, "Подію не знайдено.")
            return
        reg = row_to_dict(reg)
        event = row_to_dict(event)
        if reg.get("payment_status") in ("paid", "free"):
            await send_message(chat_id, "Ви вже оплатили.")
            return
        pay_url, _, amount = await _build_event_pay_url(dict(user), event, reg)
        role = (user.get("role") or "guest").strip()
        base_price, discount_price = _event_prices_for_role(event, role)
        pay_msg = f"💰 Оплата внеску: {amount} грн\n\n↓ Натисніть для оплати:"
        if discount_price is not None:
            pay_msg = (
                f"💰 Ціна зі знижкою: {amount} грн\n"
                f"💸 Стара ціна: {_strike_text(str(base_price))} грн\n\n"
                "↓ Натисніть для оплати:"
            )
        if pay_url:
            btns = [[{"text": "💳 Оплатити внесок", "url": pay_url}]] + back_button("events")
        else:
            pay_msg = "На жаль, не вдалося створити рахунок. Спробуйте пізніше."
            btns = back_button("events")
        await _edit_or_send_text(chat_id, message_id, pay_msg, inline_keyboard(btns))
    elif data.startswith("ev_"):
        event_id = int(data[3:])
        with db() as conn:
            event = conn.execute("SELECT * FROM events WHERE id = ? AND status = 'upcoming'", (event_id,)).fetchone()
            if not event:
                await send_message(chat_id, "Подію не знайдено.")
                return
            event = row_to_dict(event)
            role = user["role"]
            if event.get("event_type") == "closed" and role != "member":
                await send_message(chat_id, "Цей захід лише для Учасників. Бажаєте стати учасником?", reply_markup=inline_keyboard([[{"text": "Стати учасником", "callback_data": "apply_member"}]]))
                return
            lim = int(event.get("seats_limit") or 0)
            if lim:
                seat_sql = _sql_registration_occupies_seat()
                occ = conn.execute(
                    f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
                    (event_id,),
                ).fetchone()["c"]
                mine = conn.execute(
                    f"SELECT 1 FROM registrations WHERE user_id = ? AND event_id = ? AND {seat_sql}",
                    (user["id"], event_id),
                ).fetchone()
                if occ >= lim and not mine:
                    wl_en = int(event.get("waitlist_enabled") if event.get("waitlist_enabled") is not None else 1)
                    if wl_en:
                        await send_message(
                            chat_id,
                            "На жаль, місць уже немає. Можете записатися в список очікування.",
                            reply_markup=inline_keyboard([[{"text": "Список очікування", "callback_data": f"wljoin_{event_id}"}]]),
                        )
                    else:
                        await send_message(chat_id, "На жаль, для цієї події вже немає вільних місць.")
                    return
            base_price, discount_price = _event_prices_for_role(event, role)
            amount = discount_price if discount_price is not None else base_price
            now = utcnow_iso()
            conn.execute(
                """INSERT OR REPLACE INTO registrations (user_id, event_id, status, amount_uah, payment_status, created_at, updated_at)
                VALUES (?, ?, 'pending_payment', ?, 'none', ?, ?)""",
                (user["id"], event_id, amount, now, now),
            )
        if amount == 0:
            with db() as conn:
                conn.execute("UPDATE registrations SET status = 'registered', payment_status = 'free', updated_at = ? WHERE user_id = ? AND event_id = ?", (utcnow_iso(), user["id"], event_id))
            await send_message(chat_id, "Ви зареєстровані на подію.")
            schedule_event_reminder(user["id"], event_id)
            audit_log("register_free", user_id=user["id"], event_id=event_id)
            return
        order_ref = f"ev{event_id}_u{user['id']}_{int(datetime.now().timestamp())}"
        with db() as conn:
            conn.execute("UPDATE registrations SET wayforpay_order_reference = ?, updated_at = ? WHERE user_id = ? AND event_id = ?", (order_ref, utcnow_iso(), user["id"], event_id))
        u = dict(user)
        invoice_url = await create_invoice(order_ref, float(amount), f"Квиток: {event.get('title','')}", client_email=u.get("email") or "", client_phone=u.get("phone") or "", client_name=u.get("full_name") or "")
        tpl = event.get("wayforpay_member_url") if role == "member" else event.get("wayforpay_guest_url")
        pay_url = _event_payment_url_from_template(
            tpl or "",
            order_ref=order_ref,
            amount=int(amount or 0),
            event_id=event_id,
            user_id=int(user["id"]),
        ) or invoice_url or ""
        pay_msg = f"💰 Оплата внеску: {amount} грн\n\n↓ Натисніть для оплати:"
        if discount_price is not None:
            pay_msg = (
                f"💰 Ціна зі знижкою: {amount} грн\n"
                f"💸 Стара ціна: {_strike_text(str(base_price))} грн\n\n"
                "↓ Натисніть для оплати:"
            )
        if pay_url:
            btns = [[{"text": "💳 Оплатити внесок", "url": pay_url}]] + back_button("events")
        else:
            pay_msg = "На жаль, зараз неможливо створити платіжне посилання. Спробуйте пізніше або зверніться до адміна."
            btns = back_button("events")
        await _edit_or_send_text(chat_id, message_id, pay_msg, inline_keyboard(btns))
        schedule_reminder(user["id"], "stage_4_unpaid_1h", event_id=event_id)
        schedule_reminder(user["id"], "stage_4_unpaid_24h", event_id=event_id)
        audit_log("register_pending_payment", user_id=user["id"], event_id=event_id)
    elif data == "apply_member":
        role_apply = (user.get("role") or "guest").strip()
        if role_apply not in ("guest", "lead"):
            await send_message(chat_id, "Ця опція доступна для гостей та лідів.")
        else:
            with db() as conn:
                existing_app = conn.execute(
                    "SELECT 1 FROM member_applications WHERE user_id = ? AND status = 'pending'", (user["id"],)
                ).fetchone()
            if existing_app:
                await _send_system_notice(chat_id, "Ви вже подали анкету. Очікуйте зв'язку від адміна.")
            else:
                with db() as conn:
                    conn.execute(
                        "INSERT OR REPLACE INTO member_application_pending (user_id, stage) VALUES (?, 'motivation')",
                        (user["id"],),
                    )
                msg = "Розкажіть, чому ви хочете стати учасником спільноти:"
                track_id: int | None = None
                if message_id:
                    ok = await edit_message_text(chat_id, message_id, msg, reply_markup=None)
                    if ok:
                        track_id = message_id
                    else:
                        nid = await send_message(chat_id, msg, return_message_id=True)
                        track_id = int(nid) if isinstance(nid, int) else None
                else:
                    nid = await send_message(chat_id, msg, return_message_id=True)
                    track_id = int(nid) if isinstance(nid, int) else None
                if track_id:
                    _application_msg_ids[chat_id] = [track_id]
    elif data == "noop":
        pass
    elif data == "pay_membership":
        amount = 1
        period = (user.get("membership_period") or "month").strip() or "month"
        order_ref = f"mem_u{user['id']}_{int(datetime.now().timestamp())}"
        with db() as conn:
            conn.execute(
                "INSERT INTO membership_payments (user_id, wayforpay_order_reference, amount_uah, period, status, created_at) VALUES (?, ?, ?, ?, 'pending', ?)",
                (user["id"], order_ref, amount, period, utcnow_iso()),
            )
        u_dict = dict(user)
        invoice_url = await create_invoice(
            order_ref, float(amount), "Внесок учасника",
            client_email=u_dict.get("email") or "",
            client_phone=u_dict.get("phone") or "",
            client_name=u_dict.get("full_name") or "",
        )
        custom_tpl = (u_dict.get("membership_payment_url") or _get_app_config("wayforpay_membership_url", "")).strip()
        custom_url = _membership_payment_url_from_template(
            custom_tpl,
            order_ref=order_ref,
            amount=int(amount or 0),
            user_id=int(user["id"]),
            period=period,
        )
        pay_url = custom_url or invoice_url
        pay_msg = f"💰 Оплата внеску: {amount} грн\n\n↓ Натисніть для оплати:"
        if pay_url:
            btns = [[{"text": "💳 Оплатити внесок", "url": pay_url}]] + back_button("main")
        else:
            pay_msg = "На жаль, не вдалося створити рахунок. Зверніться до адміна."
            btns = back_button("main")
        await _edit_or_send_text(chat_id, message_id, pay_msg, inline_keyboard(btns))
    elif data.startswith("addtag_"):
        if user["role"] != "member":
            return
        tag_id = int(data[7:])
        with db() as conn:
            cnt = conn.execute("SELECT COUNT(*) as c FROM user_tags WHERE user_id = ?", (user["id"],)).fetchone()["c"]
            if cnt >= 10:
                msg, mk = _build_mycard_inline_view(user)
                warn = "\n\n⚠️ Максимум 10 тегів. Видаліть один, щоб додати інший."
                await _edit_or_send_text(chat_id, message_id, (msg + warn)[:4090], mk)
                return
            conn.execute("INSERT OR IGNORE INTO user_tags (user_id, tag_id) VALUES (?, ?)", (user["id"], tag_id))
        msg, mk = _build_mycard_inline_view(user)
        await _edit_or_send_text(chat_id, message_id, msg, mk)
    elif data.startswith("tag_"):
        tag_id = int(data[4:])
        with db() as conn:
            tag_name = conn.execute("SELECT name FROM member_tags WHERE id = ?", (tag_id,)).fetchone()
            members = conn.execute(
                """SELECT u.id, u.full_name, u.tg_username FROM users u
                JOIN user_tags ut ON ut.user_id = u.id WHERE ut.tag_id = ? AND u.role = 'member'""",
                (tag_id,),
            ).fetchall()
        if not members:
            msg = "Немає учасників з цим тегом."
            btns = [[{"text": "↩️ Назад", "callback_data": "back_cards"}]]
            await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
            return
        msg = f"👥 Учасники з тегом «{tag_name['name'] if tag_name else ''}»\n\n↓ Оберіть картку:"
        btns = [[{"text": m["full_name"] or "Без імені", "callback_data": f"card_{m['id']}"}] for m in members]
        btns += [[{"text": "↩️ Назад", "callback_data": "back_cards"}]]
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))
    elif data.startswith("card_"):
        member_id = int(data[5:])
        with db() as conn:
            m = conn.execute("SELECT u.*, GROUP_CONCAT(t.name) as tag_names FROM users u LEFT JOIN user_tags ut ON ut.user_id = u.id LEFT JOIN member_tags t ON t.id = ut.tag_id WHERE u.id = ? AND u.role = 'member' GROUP BY u.id", (member_id,)).fetchone()
        if not m:
            await send_message(chat_id, "Картку не знайдено.")
            return
        tg_link = f"https://t.me/{m['tg_username']}" if m.get("tg_username") else f"tg: {m['tg_user_id']}"
        msg = f"👤 {m['full_name'] or 'Без імені'}\n\n🔗 {tg_link}\n🏷 Теги: {m.get('tag_names') or 'немає'}"
        btns = [[{"text": "↩️ Назад", "callback_data": "back_cards"}]]
        await _edit_or_send_text(chat_id, message_id, msg, inline_keyboard(btns))


async def handle_contact(tg_user_id: int, contact: dict):
    chat_id = tg_user_id
    user = get_or_create_user(tg_user_id)
    phone = contact.get("phone_number")
    with db() as conn:
        conn.execute("UPDATE users SET phone = ?, updated_at = ? WHERE id = ?", (phone, utcnow_iso(), user["id"]))
        found_member = conn.execute(
            "SELECT 1 FROM users WHERE phone = ? AND role = 'member' AND id != ? LIMIT 1",
            (phone, user["id"]),
        ).fetchone()
        role = "member" if found_member else "guest"
        conn.execute("UPDATE users SET role = ?, onboarding_stage = 'done' WHERE id = ?", (role, user["id"]))
        updated = conn.execute("SELECT * FROM users WHERE id = ?", (user["id"],)).fetchone()
    updated = row_to_dict(updated)

    # Видаляємо попередні повідомлення бота (ПІБ, кнопка телефону)
    await _delete_onboarding_msgs(chat_id)

    # SendPulse у фоні, щоб не блокувати головне меню
    asyncio.create_task(sendpulse_upsert({
        "external_id": str(updated["id"]),
        "full_name": updated["full_name"] or "",
        "phone": updated["phone"] or "",
        "email": updated["email"] or "",
        "role": updated["role"],
        "telegram_id": updated["tg_user_id"],
        "tg_username": updated.get("tg_username") or "",
        "payment_status": "",
        "event_ids": "",
    }))

    role_label = "Учасник" if updated["role"] == "member" else "Гість"
    msg = f"🏠 Головне меню\n\n✅ Реєстрацію завершено.\n📋 Ваш статус: {role_label}\n\n↓ Оберіть дію:"
    await send_message(chat_id, msg, reply_markup=inline_keyboard(main_inline_menu(updated["role"])))
    await send_message(chat_id, "\u200b", reply_markup=remove_keyboard())
    schedule_reminder(updated["id"], "stage_2")
    pending_ev = updated.get("pending_deeplink_event_id")
    if pending_ev:
        uid = int(updated["id"])
        with db() as conn:
            conn.execute("UPDATE users SET pending_deeplink_event_id = NULL WHERE id = ?", (uid,))
            fresh = conn.execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
        if fresh:
            await render_event_detail(chat_id, row_to_dict(fresh), int(pending_ev))


def parse_callback(payload: dict):
    cq = payload.get("callback_query") or {}
    msg = cq.get("message") or {}
    return cq.get("id"), cq.get("from", {}).get("id"), cq.get("data"), msg.get("chat", {}).get("id"), msg.get("message_id")


def _tg_username_from_payload(payload: dict) -> str | None:
    """Витягує username з message або callback_query."""
    for key in ("message", "edited_message", "callback_query"):
        obj = payload.get(key) or {}
        frm = obj.get("from") if isinstance(obj, dict) else {}
        if frm and isinstance(frm, dict):
            un = frm.get("username")
            if un:
                return str(un).strip()
    return None


async def process_update(payload: dict) -> None:
    """Обробка оновлення від Telegram (webhook або polling)."""
    parts = parse_callback(payload)
    cq_id, cq_user_id, callback_data, chat_id, message_id = parts if len(parts) == 5 else (*parts, None)
    if cq_id and callback_data:
        username = _tg_username_from_payload(payload)
        await handle_callback(cq_id, cq_user_id, chat_id, callback_data, tg_username=username, message_id=message_id)
        return
    user, text, contact, msg_id = parse_update(payload)
    tg_user_id = user.get("id") if isinstance(user, dict) else (user["id"] if user else None)
    if not tg_user_id:
        return
    tg_username = (user.get("username") or "").strip() or None
    get_or_create_user(tg_user_id, tg_username=tg_username)
    if text and (text == "/start" or text.startswith("/start ")):
        start_arg = text.split(maxsplit=1)[1] if " " in text.strip() else None
        await handle_start(tg_user_id, start_arg=start_arg)
    elif contact:
        await handle_contact(tg_user_id, contact)
    elif text:
        await handle_text(tg_user_id, text.strip() if text else "", input_message_id=msg_id)


async def polling_worker() -> None:
    """Long polling — працює без домену та HTTPS."""
    if not settings.telegram_bot_token:
        print("[BOT] Помилка: TELEGRAM_BOT_TOKEN не вказано в .env")
        return
    try:
        await delete_webhook()
        print("[BOT] Webhook видалено, polling запущено")
    except Exception as e:
        print(f"[BOT] deleteWebhook: {e}")
    offset = None
    while True:
        try:
            updates = await get_updates(offset=offset)
            for u in updates:
                offset = u.get("update_id", 0) + 1
                try:
                    await process_update(u)
                except Exception as e:
                    import traceback
                    print(f"[BOT] Помилка: {e}")
                    traceback.print_exc()
        except Exception as e:
            print(f"[BOT] getUpdates: {e}")
        await asyncio.sleep(0.5)


async def payment_poll_worker() -> None:
    """Перевірка статусу оплати WayForPay без callback (для роботи без домену)."""
    while True:
        try:
            with db() as conn:
                rows = conn.execute(
                    """
                    SELECT r.id, r.user_id, r.event_id, r.wayforpay_order_reference, r.payment_status
                    FROM registrations r
                    WHERE r.wayforpay_order_reference IS NOT NULL
                      AND r.payment_status NOT IN ('paid', 'free')
                      AND r.created_at > datetime('now', '-2 days')
                    LIMIT 20
                    """
                ).fetchall()
            for row in rows:
                status = await check_payment_status(row["wayforpay_order_reference"])
                if status == "Approved":
                    event_title = ""
                    u = None
                    with db() as conn:
                        conn.execute(
                            "UPDATE registrations SET payment_status='paid', status='registered', updated_at=? WHERE id=?",
                            (utcnow_iso(), row["id"]),
                        )
                        e = conn.execute("SELECT title FROM events WHERE id = ?", (row["event_id"],)).fetchone()
                        if e:
                            event_title = (e["title"] or "").strip()
                        u = conn.execute("SELECT * FROM users WHERE id=?", (row["user_id"],)).fetchone()
                        if u and (u["role"] or "").strip() == "lead":
                            conn.execute("UPDATE users SET role='guest', updated_at=? WHERE id=?", (utcnow_iso(), row["user_id"]))
                    audit_log("payment_event", user_id=row["user_id"], event_id=row["event_id"], detail=row["wayforpay_order_reference"] or "")
                    if u and u["tg_user_id"]:
                        title_part = f" «{event_title}»" if event_title else ""
                        await send_message(
                            u["tg_user_id"],
                            f"✅ Ви успішно оплатили участь у події{title_part}.",
                            reply_markup=inline_keyboard([[{"text": "Оновити статус 🔄", "callback_data": "refresh_member_menu"}]]),
                        )
                    schedule_event_reminder(row["user_id"], row["event_id"])
                    if u:
                        u = dict(u)
                        if (u.get("role") or "").strip() == "lead":
                            u["role"] = "guest"
                        await sendpulse_upsert({
                    "external_id": str(u["id"]), "full_name": u.get("full_name") or "",
                    "phone": u.get("phone") or "", "email": u.get("email") or "",
                    "role": u.get("role") or "guest", "telegram_id": u.get("tg_user_id"),
                    "tg_username": u.get("tg_username") or "", "payment_status": "paid",
                    "event_ids": str(row["event_id"]),
                })
            with db() as conn:
                mps = conn.execute(
                    "SELECT * FROM membership_payments WHERE wayforpay_order_reference IS NOT NULL AND status = 'pending' AND created_at > datetime('now', '-2 days') LIMIT 10"
                ).fetchall()
            for mp in mps:
                status = await check_payment_status(mp["wayforpay_order_reference"])
                if status == "Approved":
                    with db() as conn:
                        conn.execute("UPDATE membership_payments SET status = 'paid' WHERE id = ?", (mp["id"],))
                        _extend_membership(mp["user_id"], mp["period"])
                        user_row = conn.execute("SELECT tg_user_id FROM users WHERE id = ?", (mp["user_id"],)).fetchone()
                    if user_row:
                        with db() as c2:
                            u2 = c2.execute("SELECT * FROM users WHERE id = ?", (mp["user_id"],)).fetchone()
                        if u2:
                            u2 = dict(u2)
                            period_name = {"month": "місяць", "quarter": "квартал", "year": "рік"}.get(u2.get("membership_period") or "month", "місяць")
                            exp = (u2.get("membership_expires_at") or "")[:10]
                            await send_message(user_row["tg_user_id"], f"✅ Оплата успішна! Ваше членство продовжено.\n\n📅 Термін до: {exp or '—'}\n💰 Наступний внесок: {(u2.get('membership_amount_uah') or 0)} грн ({period_name})")
        except Exception as e:
            print(f"[PAY] Помилка: {e}")
        await asyncio.sleep(60)


async def event_archive_worker():
    """Архів подій після старту: no_show, автозапит фідбеку для check-in."""
    while True:
        try:
            now_iso = datetime.now(timezone.utc).isoformat()
            archived_ids: list[int] = []
            feedback_targets: list[tuple[int, int]] = []
            with db() as conn:
                rows = conn.execute(
                    "SELECT id FROM events WHERE status = 'upcoming' AND starts_at <= ?",
                    (now_iso,),
                ).fetchall()
                archived_ids = [int(r["id"]) for r in rows]
                conn.execute("UPDATE events SET status = 'archived' WHERE status = 'upcoming' AND starts_at <= ?", (now_iso,))
                for eid in archived_ids:
                    conn.execute(
                        """UPDATE registrations SET no_show = 1, updated_at = ?
                        WHERE event_id = ? AND attended = 0 AND status = 'registered'
                        AND payment_status IN ('paid', 'free')""",
                        (utcnow_iso(), eid),
                    )
                    regs = conn.execute(
                        """SELECT r.user_id, u.tg_user_id FROM registrations r
                        JOIN users u ON u.id = r.user_id
                        WHERE r.event_id = ? AND r.attended = 1 AND r.feedback_sent = 0""",
                        (eid,),
                    ).fetchall()
                    for reg in regs:
                        conn.execute(
                            """INSERT OR REPLACE INTO feedback_pending (user_id, event_id, stage, created_at, nudge_sent)
                            VALUES (?, ?, 'waiting_rating', ?, 0)""",
                            (reg["user_id"], eid, utcnow_iso()),
                        )
                        feedback_targets.append((int(reg["tg_user_id"]), eid))
            for tid, _eid in feedback_targets:
                try:
                    await send_message(
                        tid,
                        "Дякуємо, що були на заході! Оцініть подію від 1 до 5 (напишіть цифру).",
                    )
                except Exception:
                    pass
        except Exception:
            pass
        await asyncio.sleep(300)


async def feedback_nudge_worker():
    """Повторне нагадування про оцінку (ТЗ §5.2), якщо за добу не оцінили."""
    while True:
        try:
            with db() as conn:
                rows = conn.execute(
                    """SELECT fp.user_id, u.tg_user_id FROM feedback_pending fp
                    JOIN users u ON u.id = fp.user_id
                    WHERE fp.stage = 'waiting_rating' AND IFNULL(fp.nudge_sent, 0) = 0
                    AND fp.created_at IS NOT NULL
                    AND datetime(fp.created_at) <= datetime('now', '-1 day')"""
                ).fetchall()
            for row in rows:
                try:
                    await send_message(
                        int(row["tg_user_id"]),
                        "Нагадування: будь ласка, оцініть подію цифрою від 1 до 5.",
                    )
                except Exception:
                    pass
                with db() as conn:
                    conn.execute("UPDATE feedback_pending SET nudge_sent = 1 WHERE user_id = ?", (row["user_id"],))
        except Exception:
            pass
        await asyncio.sleep(600)


@app.on_event("startup")
async def startup_event():
    init_db()
    asyncio.create_task(reminder_worker())
    asyncio.create_task(event_timeline_worker())
    asyncio.create_task(feedback_nudge_worker())
    asyncio.create_task(payment_poll_worker())
    asyncio.create_task(event_archive_worker())
    if settings.use_polling:
        asyncio.create_task(polling_worker())
        print("[BOT] Запущено в режимі polling (без домену). Чекайте повідомлення...")


@app.get("/health")
async def health():
    return {"ok": True}


@app.post("/telegram/webhook")
async def telegram_webhook(request: Request):
    if not settings.use_polling:
        if request.headers.get("x-telegram-bot-api-secret-token") != settings.telegram_webhook_secret:
            raise HTTPException(status_code=403, detail="invalid secret")
    payload = await request.json()
    try:
        await process_update(payload)
    except Exception as e:
        import traceback

        print(f"[BOT] webhook process_update: {e}")
        traceback.print_exc()
    return {"ok": True}


def _extend_membership(user_id: int, period: str) -> None:
    """Продовжує членство на період."""
    months = {"month": 1, "quarter": 3, "year": 12}.get(period, 1)
    with db() as conn:
        u = conn.execute("SELECT membership_expires_at FROM users WHERE id = ?", (user_id,)).fetchone()
        base = datetime.now(timezone.utc)
        if u and u["membership_expires_at"]:
            try:
                base = datetime.fromisoformat(u["membership_expires_at"].replace("Z", "+00:00"))
                if base < datetime.now(timezone.utc):
                    base = datetime.now(timezone.utc)
            except Exception:
                pass
        new_exp = (base + timedelta(days=30 * months)).isoformat()
        conn.execute("UPDATE users SET membership_expires_at = ?, updated_at = ? WHERE id = ?", (new_exp, utcnow_iso(), user_id))


@app.post("/wayforpay/callback")
async def wayforpay_callback(request: Request):
    try:
        payload = await request.json()
    except Exception:
        payload = dict(await request.form())
    if not verify_callback(payload):
        raise HTTPException(status_code=403, detail="Invalid signature")
    order_ref = payload.get("orderReference", "")
    status = payload.get("transactionStatus", "")
    if order_ref.startswith("mem_"):
        u2 = None
        with db() as conn:
            mp = conn.execute("SELECT * FROM membership_payments WHERE wayforpay_order_reference = ?", (order_ref,)).fetchone()
            if not mp:
                raise HTTPException(status_code=404, detail="membership payment not found")
            if status == "Approved":
                conn.execute("UPDATE membership_payments SET status = 'paid' WHERE id = ?", (mp["id"],))
                _extend_membership(mp["user_id"], mp["period"])
                user = conn.execute("SELECT * FROM users WHERE id = ?", (mp["user_id"],)).fetchone()
                u2 = dict(user) if user else None
        if status == "Approved" and u2:
            period_name = {"month": "місяць", "quarter": "квартал", "year": "рік"}.get(u2.get("membership_period") or "month", "місяць")
            exp = (u2.get("membership_expires_at") or "")[:10]
            await send_message(u2["tg_user_id"], f"✅ Оплата успішна! Ваше членство продовжено.\n\n📅 Термін до: {exp or '—'}\n💰 Наступний внесок: {(u2.get('membership_amount_uah') or 0)} грн ({period_name})")
            await sendpulse_upsert({
                "external_id": str(u2["id"]), "full_name": u2.get("full_name") or "",
                "phone": u2.get("phone") or "", "email": u2.get("email") or "",
                "role": u2.get("role") or "member", "telegram_id": u2.get("tg_user_id"),
                "tg_username": u2.get("tg_username") or "", "payment_status": "approved",
                "event_ids": "",
            })
    else:
        event_title = ""
        with db() as conn:
            reg = conn.execute("SELECT * FROM registrations WHERE wayforpay_order_reference = ?", (order_ref,)).fetchone()
            if not reg:
                raise HTTPException(status_code=404, detail="registration not found")
            if status == "Approved":
                conn.execute(
                    "UPDATE registrations SET payment_status = 'paid', status = 'registered', updated_at = ? WHERE id = ?",
                    (utcnow_iso(), reg["id"]),
                )
                e = conn.execute("SELECT title FROM events WHERE id = ?", (reg["event_id"],)).fetchone()
                if e:
                    event_title = (e["title"] or "").strip()
            user = conn.execute("SELECT * FROM users WHERE id = ?", (reg["user_id"],)).fetchone()
            if user and status == "Approved" and ((user["role"] or "").strip() == "lead"):
                conn.execute("UPDATE users SET role='guest', updated_at=? WHERE id=?", (utcnow_iso(), reg["user_id"]))
        if user and status == "Approved":
            user = dict(user)
            if status == "Approved" and (user.get("role") or "").strip() == "lead":
                user["role"] = "guest"
            title_part = f" «{event_title}»" if event_title else ""
            await send_message(
                user["tg_user_id"],
                f"✅ Ви успішно оплатили участь у події{title_part}.",
                reply_markup=inline_keyboard([[{"text": "Оновити статус 🔄", "callback_data": "refresh_member_menu"}]]),
            )
            audit_log("payment_event", user_id=reg["user_id"], event_id=reg["event_id"], detail=order_ref)
            schedule_event_reminder(reg["user_id"], reg["event_id"])
            await sendpulse_upsert({
                "external_id": str(user["id"]), "full_name": user.get("full_name") or "",
                "phone": user.get("phone") or "", "email": user.get("email") or "",
                "role": user.get("role") or "guest", "telegram_id": user.get("tg_user_id"),
                "tg_username": user.get("tg_username") or "", "payment_status": "paid",
                "event_ids": str(reg["event_id"]),
            })
    return JSONResponse(success_ack(order_ref))


@app.get("/admin/entry/{token}")
async def admin_entry(token: str):
    d = _decode_signed_token(token)
    if not d or "tg_user_id" not in d:
        raise HTTPException(status_code=403, detail="Invalid or expired token")
    return RedirectResponse(url=f"/admin?token={token}")


@app.get("/admin", response_class=HTMLResponse)
async def admin_page(request: Request, token: str = ""):
    if not token:
        return templates.TemplateResponse("admin_login.html", {"request": request})
    d = _decode_signed_token(token)
    if not d or "tg_user_id" not in d:
        return templates.TemplateResponse("admin_login.html", {"request": request, "error": "Невалідний токен"})
    with db() as conn:
        try:
            apps_count = conn.execute("SELECT COUNT(*) as c FROM member_applications WHERE status='pending'").fetchone()["c"]
        except Exception:
            apps_count = 0
        stats = {
            "users": conn.execute("SELECT COUNT(*) as c FROM users").fetchone()["c"],
            "members": conn.execute("SELECT COUNT(*) as c FROM users WHERE role='member'").fetchone()["c"],
            "guests": conn.execute("SELECT COUNT(*) as c FROM users WHERE role='guest'").fetchone()["c"],
            "leads": conn.execute("SELECT COUNT(*) as c FROM users WHERE role='lead'").fetchone()["c"],
            "events": conn.execute("SELECT COUNT(*) as c FROM events").fetchone()["c"],
            "registrations": conn.execute("SELECT COUNT(*) as c FROM registrations").fetchone()["c"],
            "applications": apps_count,
        }
        events = conn.execute("SELECT * FROM events WHERE status='upcoming' ORDER BY starts_at ASC LIMIT 50").fetchall()
        archived = conn.execute("SELECT * FROM events WHERE status='archived' ORDER BY starts_at DESC LIMIT 20").fetchall()
        users = conn.execute("SELECT * FROM users ORDER BY id DESC LIMIT 50").fetchall()
    return templates.TemplateResponse(
        "admin_dashboard.html", {"request": request, "token": token, "stats": stats, "events": events, "archived": archived, "users": users}
    )


@app.post("/admin/events/create")
async def admin_create_event(
    token: str = Form(...),
    title: str = Form(...),
    starts_at: str = Form(...),
    guest_price_uah: int = Form(2),
    lead_price_uah: int = Form(2),
    member_price_uah: int = Form(1),
    guest_discount_price_uah: int | None = Form(None),
    lead_discount_price_uah: int | None = Form(None),
    member_discount_price_uah: int | None = Form(None),
    event_type: str = Form("open"),
    description: str = Form(""),
    location: str = Form(""),
    photo_file: UploadFile = File(None),
    video_file: UploadFile = File(None),
    zoom_url: str = Form(""),
    seats_limit: int = Form(0),
    wayforpay_guest_url: str = Form(""),
    wayforpay_member_url: str = Form(""),
    wayforpay_lead_url: str = Form(""),
    msg_rsvp_40h: str = Form(""),
    msg_remind_24h: str = Form(""),
    msg_remind_3h: str = Form(""),
    waitlist_enabled: str = Form("1"),
):
    _verify_admin_token(token)
    banner_url = None
    if photo_file and photo_file.filename:
        data = await photo_file.read()
        if data and len(data) < 10 * 1024 * 1024:
            ext = Path(photo_file.filename or "img").suffix or ".jpg"
            fname = f"evp_new_{int(datetime.now(timezone.utc).timestamp())}{ext}"
            (UPLOADS_DIR / fname).write_bytes(data)
            banner_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
    elif video_file and video_file.filename:
        data = await video_file.read()
        if data and len(data) < 50 * 1024 * 1024:
            ext = Path(video_file.filename or "vid").suffix or ".mp4"
            fname = f"evv_new_{int(datetime.now(timezone.utc).timestamp())}{ext}"
            (UPLOADS_DIR / fname).write_bytes(data)
            banner_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
    wl = 1 if (waitlist_enabled or "").strip() == "1" else 0
    with db() as conn:
        conn.execute(
            """
            INSERT INTO events (title, starts_at, guest_price_uah, lead_price_uah, member_price_uah, guest_discount_price_uah, lead_discount_price_uah, member_discount_price_uah, event_type, description, location, banner_url, zoom_url, seats_limit, wayforpay_guest_url, wayforpay_member_url, wayforpay_lead_url, msg_rsvp_40h, msg_remind_24h, msg_remind_3h, waitlist_enabled)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                title,
                starts_at,
                guest_price_uah,
                lead_price_uah,
                member_price_uah,
                guest_discount_price_uah,
                lead_discount_price_uah,
                member_discount_price_uah,
                event_type,
                description,
                location,
                banner_url,
                zoom_url,
                seats_limit,
                (wayforpay_guest_url or "").strip() or None,
                (wayforpay_member_url or "").strip() or None,
                (wayforpay_lead_url or "").strip() or None,
                (msg_rsvp_40h or "").strip() or None,
                (msg_remind_24h or "").strip() or None,
                (msg_remind_3h or "").strip() or None,
                wl,
            ),
        )
    return RedirectResponse(url=f"/admin?token={token}", status_code=303)


@app.post("/admin/events/{event_id}/update")
async def admin_update_event(
    event_id: int,
    token: str = Form(...),
    title: str = Form(...),
    starts_at: str = Form(...),
    guest_price_uah: int = Form(...),
    lead_price_uah: int = Form(2),
    member_price_uah: int = Form(0),
    guest_discount_price_uah: int | None = Form(None),
    lead_discount_price_uah: int | None = Form(None),
    member_discount_price_uah: int | None = Form(None),
    event_type: str = Form("open"),
    description: str = Form(""),
    location: str = Form(""),
    photo_file: UploadFile = File(None),
    video_file: UploadFile = File(None),
    zoom_url: str = Form(""),
    seats_limit: int = Form(0),
    wayforpay_guest_url: str = Form(""),
    wayforpay_member_url: str = Form(""),
    wayforpay_lead_url: str = Form(""),
    msg_rsvp_40h: str = Form(""),
    msg_remind_24h: str = Form(""),
    msg_remind_3h: str = Form(""),
    waitlist_enabled: str = Form("1"),
    push_on_change: str = Form(""),
    push_message: str = Form(""),
    media_type: str = Form("text"),
    attachment_file: UploadFile = File(None),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_admin_token(token)
    wl = 1 if (waitlist_enabled or "").strip() == "1" else 0
    with db() as conn:
        old = conn.execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
        if not old:
            raise HTTPException(status_code=404, detail="Event not found")
        old_d = row_to_dict(old)
        banner_url = old_d.get("banner_url")
        if photo_file and photo_file.filename:
            data = await photo_file.read()
            if data and len(data) < 10 * 1024 * 1024:
                ext = Path(photo_file.filename or "img").suffix or ".jpg"
                fname = f"evp_{event_id}_{int(datetime.now(timezone.utc).timestamp())}{ext}"
                (UPLOADS_DIR / fname).write_bytes(data)
                banner_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
        elif video_file and video_file.filename:
            data = await video_file.read()
            if data and len(data) < 50 * 1024 * 1024:
                ext = Path(video_file.filename or "vid").suffix or ".mp4"
                fname = f"evv_{event_id}_{int(datetime.now(timezone.utc).timestamp())}{ext}"
                (UPLOADS_DIR / fname).write_bytes(data)
                banner_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
        conn.execute(
            """UPDATE events SET title=?, starts_at=?, guest_price_uah=?, lead_price_uah=?, member_price_uah=?, guest_discount_price_uah=?, lead_discount_price_uah=?, member_discount_price_uah=?, event_type=?,
            description=?, location=?, banner_url=?, zoom_url=?, seats_limit=?,
            msg_rsvp_40h=?, msg_remind_24h=?, msg_remind_3h=?, waitlist_enabled=?,
            wayforpay_guest_url=?, wayforpay_member_url=?, wayforpay_lead_url=? WHERE id=?""",
            (
                title,
                starts_at,
                guest_price_uah,
                lead_price_uah,
                member_price_uah,
                guest_discount_price_uah,
                lead_discount_price_uah,
                member_discount_price_uah,
                event_type,
                description,
                location,
                banner_url or None,
                zoom_url or None,
                seats_limit,
                (msg_rsvp_40h or "").strip() or None,
                (msg_remind_24h or "").strip() or None,
                (msg_remind_3h or "").strip() or None,
                wl,
                (wayforpay_guest_url or "").strip() or None,
                (wayforpay_member_url or "").strip() or None,
                (wayforpay_lead_url or "").strip() or None,
                event_id,
            ),
        )
        if push_on_change == "1":
            msg_text = (push_message or "").strip()
            uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
            if not uploads and attachment_file and attachment_file.filename:
                uploads = [attachment_file]
            media_items: list[tuple[str, str]] = []
            if uploads:
                for up in uploads:
                    media_url, content_type = await _save_uploaded_media(up, f"evupd_{event_id}")
                    if media_url:
                        media_items.append((media_url, content_type))
            if msg_text or media_items:
                rows = conn.execute(
                    "SELECT u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id WHERE r.event_id = ? AND r.status IN ('registered', 'pending_payment')",
                    (event_id,),
                ).fetchall()
                for row in rows:
                    rd = row_to_dict(row)
                    tid = rd.get("tg_user_id")
                    if tid:
                        if media_items:
                            first = True
                            for m_url, m_type in media_items:
                                cap = msg_text[:1024] if first else ""
                                if m_type == "photo":
                                    await send_photo(tid, m_url, caption=cap or None)
                                elif m_type == "video":
                                    await send_video(tid, m_url, caption=cap or None)
                                else:
                                    await send_document(tid, m_url, caption=cap or None)
                                first = False
                        elif msg_text:
                            await send_message(tid, msg_text)
    return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)


@app.get("/admin/events/{event_id}", response_class=HTMLResponse)
async def admin_event_detail(request: Request, event_id: int, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        event = conn.execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
        if not event:
            raise HTTPException(status_code=404, detail="Event not found")
        regs = conn.execute(
            """
            SELECT r.*, u.id as user_id, u.full_name, u.phone, u.email, u.tg_user_id
            FROM registrations r
            JOIN users u ON u.id = r.user_id
            WHERE r.event_id = ? AND r.status IN ('registered', 'pending_payment')
            ORDER BY r.attended DESC, u.full_name
            """,
            (event_id,),
        ).fetchall()
        feedbacks = conn.execute(
            "SELECT f.*, u.full_name FROM feedback f JOIN users u ON u.id = f.user_id WHERE f.event_id = ? ORDER BY f.created_at DESC",
            (event_id,),
        ).fetchall()
        seat_sql = _sql_registration_occupies_seat()
        seg_all = conn.execute("SELECT COUNT(*) as c FROM registrations WHERE event_id = ?", (event_id,)).fetchone()["c"]
        seg_paid = conn.execute(
            "SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND status = 'registered' AND payment_status IN ('paid','free')",
            (event_id,),
        ).fetchone()["c"]
        seg_unpaid = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ?
            AND NOT (status = 'registered' AND payment_status IN ('paid','free'))""",
            (event_id,),
        ).fetchone()["c"]
        seg_conf = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND status = 'registered'
            AND payment_status IN ('paid','free') AND IFNULL(rsvp_will_attend,0) = 1""",
            (event_id,),
        ).fetchone()["c"]
        seg_noconf = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND status = 'registered'
            AND payment_status IN ('paid','free') AND rsvp_will_attend IS NULL""",
            (event_id,),
        ).fetchone()["c"]
        occ = conn.execute(
            f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
            (event_id,),
        ).fetchone()["c"]
    event_d = row_to_dict(event) if event else {}
    reg_rows = []
    for r in regs:
        d = row_to_dict(r)
        d["tz_status"] = registration_tz_label(d, event_d)
        reg_rows.append(d)
    lim = int(event_d.get("seats_limit") or 0)
    seats_free = max(0, lim - occ) if lim else None
    bot_u = (_get_app_config("telegram_bot_username", "") or "").strip().lstrip("@")
    event_deeplink = f"https://t.me/{bot_u}?start=e{event_id}" if bot_u else ""
    return templates.TemplateResponse(
        "admin_event.html",
        {
            "request": request,
            "token": token,
            "event": event_d,
            "registrations": reg_rows,
            "feedbacks": [row_to_dict(f) for f in feedbacks],
            "event_deeplink": event_deeplink,
            "seg": {"all": seg_all, "paid": seg_paid, "unpaid": seg_unpaid, "confirmed": seg_conf, "noconf": seg_noconf, "seats_free": seats_free},
        },
    )


@app.post("/admin/events/{event_id}/archive")
async def admin_archive_event(event_id: int, token: str = Form(...)):
    _verify_admin_token(token)
    with db() as conn:
        conn.execute("UPDATE events SET status = 'archived' WHERE id = ?", (event_id,))
    return RedirectResponse(url=f"/admin?token={token}", status_code=303)


@app.post("/admin/events/{event_id}/restore")
async def admin_restore_event(event_id: int, token: str = Form(...), starts_at: str = Form("")):
    _verify_admin_token(token)
    starts_at_clean = (starts_at or "").strip()
    with db() as conn:
        if starts_at_clean:
            conn.execute("UPDATE events SET status = 'upcoming', starts_at = ? WHERE id = ?", (starts_at_clean, event_id))
        else:
            conn.execute("UPDATE events SET status = 'upcoming' WHERE id = ?", (event_id,))
    return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)


@app.post("/admin/events/{event_id}/checkin")
async def admin_checkin(event_id: int, token: str = Form(...), reg_id: str = Form(...), attended: str = Form("1")):
    _verify_admin_token(token)
    with db() as conn:
        row = conn.execute("SELECT user_id FROM registrations WHERE id = ? AND event_id = ?", (reg_id, event_id)).fetchone()
        conn.execute("UPDATE registrations SET attended = ? WHERE id = ? AND event_id = ?", (1 if attended == "1" else 0, reg_id, event_id))
    if row:
        audit_log("checkin", user_id=row["user_id"], event_id=event_id, detail=f"attended={attended}")
    return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)


@app.post("/admin/events/{event_id}/feedback_request")
async def admin_feedback_request(event_id: int, token: str = Form(...)):
    _verify_admin_token(token)
    with db() as conn:
        rows = conn.execute(
            "SELECT r.user_id, u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id WHERE r.event_id = ? AND r.attended = 1 AND r.feedback_sent = 0",
            (event_id,),
        ).fetchall()
        for row in rows:
            conn.execute(
                "INSERT OR REPLACE INTO feedback_pending (user_id, event_id, stage, created_at, nudge_sent) VALUES (?, ?, 'waiting_rating', ?, 0)",
                (row["user_id"], event_id, utcnow_iso()),
            )
        for row in rows:
            await send_message(row["tg_user_id"], "Дякуємо, що були на заході! Оцініть подію від 1 до 5 (напишіть цифру).")
    return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)


@app.post("/admin/events/{event_id}/noshow_push")
async def admin_noshow_push(
    event_id: int,
    token: str = Form(...),
    message: str = Form("Ви не зʼявилися на заході. Чи можете пояснити причину?"),
    media_type: str = Form("text"),
    attachment_file: UploadFile = File(None),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_admin_token(token)
    text = (message or "").strip()
    uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
    if not uploads and attachment_file and attachment_file.filename:
        uploads = [attachment_file]
    media_items: list[tuple[str, str]] = []
    if uploads:
        for up in uploads:
            media_url, content_type = await _save_uploaded_media(up, f"noshow_{event_id}")
            if media_url:
                media_items.append((media_url, content_type))
    with db() as conn:
        rows = conn.execute(
            "SELECT u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id WHERE r.event_id = ? AND r.attended = 0 AND r.status = 'registered'",
            (event_id,),
        ).fetchall()
    for row in rows:
        tid = row["tg_user_id"]
        if media_items:
            first = True
            for m_url, m_type in media_items:
                cap = text[:1024] if first else ""
                if m_type == "photo":
                    await send_photo(tid, m_url, caption=cap or None)
                elif m_type == "video":
                    await send_video(tid, m_url, caption=cap or None)
                else:
                    await send_document(tid, m_url, caption=cap or None)
                first = False
        else:
            await send_message(tid, text or "Нагадування")
    return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)


@app.get("/admin/broadcast", response_class=HTMLResponse)
async def admin_broadcast_page(request: Request, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        events = conn.execute("SELECT id, title FROM events WHERE status = 'upcoming'").fetchall()
    return templates.TemplateResponse("admin_broadcast.html", {"request": request, "token": token, "events": events})


@app.post("/admin/broadcast/send")
async def admin_broadcast_send(
    token: str = Form(...),
    segment: str = Form(...),
    event_id: str = Form(""),
    message: str = Form(...),
    msg_type: str = Form("text"),
    restrict_tg_ids: str = Form(""),
    photo_file: UploadFile = File(None),
    video_file: UploadFile = File(None),
    attachment_file: UploadFile = File(None),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_admin_token(token)
    if segment in ("event_guests", "event_unpaid", "event_paid_no_rsvp") and not (event_id or "").strip():
        raise HTTPException(status_code=400, detail="Оберіть подію для цього сегменту")
    with db() as conn:
        if segment == "members":
            rows = conn.execute("SELECT tg_user_id FROM users WHERE role = 'member'").fetchall()
        elif segment == "guests":
            rows = conn.execute("SELECT tg_user_id FROM users WHERE role = 'guest'").fetchall()
        elif segment == "leads":
            rows = conn.execute("SELECT tg_user_id FROM users WHERE role = 'lead'").fetchall()
        elif segment == "event_guests" and event_id:
            rows = conn.execute(
                "SELECT u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id WHERE r.event_id = ?",
                (int(event_id),),
            ).fetchall()
        elif segment == "event_unpaid" and event_id:
            rows = conn.execute(
                """SELECT u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id
                WHERE r.event_id = ? AND NOT (r.status = 'registered' AND r.payment_status IN ('paid','free'))""",
                (int(event_id),),
            ).fetchall()
        elif segment == "event_paid_no_rsvp" and event_id:
            rows = conn.execute(
                """SELECT u.tg_user_id FROM registrations r JOIN users u ON u.id = r.user_id
                WHERE r.event_id = ? AND r.status = 'registered' AND r.payment_status IN ('paid','free')
                AND r.rsvp_will_attend IS NULL""",
                (int(event_id),),
            ).fetchall()
        else:
            rows = conn.execute("SELECT tg_user_id FROM users").fetchall()
    restrict = _parse_tg_restrict_ids(restrict_tg_ids)
    if restrict is not None:
        rows = [r for r in rows if r["tg_user_id"] in restrict]
    media_url = ""
    media_items: list[tuple[str, str]] = []
    resolved_type = (msg_type or "text").strip()
    uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
    if not uploads and attachment_file and attachment_file.filename:
        uploads = [attachment_file]
    if uploads:
        for up in uploads:
            media_url_item, detected_type = await _save_uploaded_media(up, "bc")
            if media_url_item:
                media_items.append((media_url_item, detected_type))
        if media_items:
            resolved_type = media_items[0][1]
    elif resolved_type == "photo" and photo_file and photo_file.filename:
        data = await photo_file.read()
        if data and len(data) < 10 * 1024 * 1024:
            ext = Path(photo_file.filename or "img").suffix or ".jpg"
            fname = f"bc_p_{int(datetime.now(timezone.utc).timestamp())}{ext}"
            (UPLOADS_DIR / fname).write_bytes(data)
            media_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
    elif resolved_type == "video" and video_file and video_file.filename:
        data = await video_file.read()
        if data and len(data) < 50 * 1024 * 1024:
            ext = Path(video_file.filename or "vid").suffix or ".mp4"
            fname = f"bc_v_{int(datetime.now(timezone.utc).timestamp())}{ext}"
            (UPLOADS_DIR / fname).write_bytes(data)
            media_url = f"{settings.app_base_url.rstrip('/')}/uploads/{fname}"
    caption = message[:1024] if message else ""
    seen: set[int] = set()
    for row in rows:
        tid = row["tg_user_id"]
        if not tid or tid in seen:
            continue
        seen.add(tid)
        if media_items:
            first = True
            for m_url, m_type in media_items:
                cap = caption if first else ""
                if m_type == "photo":
                    await send_photo(tid, m_url, caption=cap)
                elif m_type == "video":
                    await send_video(tid, m_url, caption=cap)
                else:
                    await send_document(tid, m_url, caption=cap)
                first = False
        elif resolved_type == "photo" and media_url:
            await send_photo(tid, media_url, caption=caption)
        elif resolved_type == "video" and media_url:
            await send_video(tid, media_url, caption=caption)
        else:
            await send_message(tid, message)
    return RedirectResponse(url=f"/admin?token={token}", status_code=303)


@app.get("/admin/applications", response_class=HTMLResponse)
async def admin_applications(request: Request, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        apps = conn.execute(
            """SELECT ma.*, u.full_name, u.phone, u.email, u.tg_user_id, u.tg_username FROM member_applications ma
            JOIN users u ON u.id = ma.user_id WHERE ma.status = 'pending' ORDER BY ma.created_at DESC"""
        ).fetchall()
    apps_data = []
    for a in apps:
        d = dict(a)
        d["tg_display"] = format_tg_display(d)
        apps_data.append(d)
    return templates.TemplateResponse(
        "admin_applications.html",
        {"request": request, "token": token, "applications": apps_data},
    )


@app.post("/admin/applications/{app_id}/approve")
async def admin_approve_application(
    app_id: int,
    token: str = Form(...),
    membership_period: str = Form("month"),
    membership_amount_uah: int = Form(100),
    payment_url: str = Form(""),
):
    _verify_admin_token(token)
    payment_url_clean = (payment_url or "").strip()
    payment_url_db = payment_url_clean if payment_url_clean.startswith("http") else None
    with db() as conn:
        app = conn.execute(
            "SELECT user_id FROM member_applications WHERE id = ? AND status = 'pending'", (app_id,)
        ).fetchone()
        if app:
            conn.execute(
                "UPDATE users SET role = 'member', membership_period = ?, membership_amount_uah = ?, membership_payment_url = ? WHERE id = ?",
                (membership_period, membership_amount_uah, payment_url_db, app["user_id"]),
            )
            conn.execute("UPDATE member_applications SET status = 'approved' WHERE id = ?", (app_id,))
            u = conn.execute("SELECT * FROM users WHERE id = ?", (app["user_id"],)).fetchone()
            if u:
                u = dict(u)
                period_name = {"month": "місяць", "quarter": "квартал", "year": "рік"}.get(membership_period, membership_period)
                if membership_amount_uah == 0:
                    _extend_membership(u["id"], membership_period)
                msg = "🎉 Вітаємо! Вам підтверджено статус Учасника спільноти.\n\n"
                if membership_amount_uah > 0:
                    if payment_url_clean.startswith("http"):
                        msg += f"💰 Ціна за участь: {membership_amount_uah} грн ({period_name}).\n\n"
                    else:
                        msg += f"💰 Ціна за участь: {membership_amount_uah} грн ({period_name}). Оплатити можна в меню «Мій кабінет» — там є кнопка «Оплатити внесок».\n\n"
                msg += "Оновіть меню, щоб отримати доступ до меню Учасника."
                btns = [[{"text": "Оновити статус 🔄", "callback_data": "refresh_member_menu"}]]
                if payment_url_clean.startswith("http"):
                    btns.insert(0, [{"text": "💳 Оплатити внесок", "url": payment_url_clean}])
                await send_message(u["tg_user_id"], msg, reply_markup=inline_keyboard(btns))
                await sendpulse_upsert({"external_id": str(u["id"]), "full_name": u.get("full_name") or "", "phone": u.get("phone") or "", "email": u.get("email") or "", "role": "member", "telegram_id": u.get("tg_user_id"), "tg_username": u.get("tg_username") or "", "payment_status": "pending" if membership_amount_uah > 0 else "approved", "event_ids": ""})
    return RedirectResponse(url=f"/admin/applications?token={token}", status_code=303)


@app.post("/admin/applications/{app_id}/reject")
async def admin_reject_application(app_id: int, token: str = Form(...)):
    _verify_admin_token(token)
    with db() as conn:
        app = conn.execute(
            "SELECT user_id FROM member_applications WHERE id = ? AND status = 'pending'", (app_id,)
        ).fetchone()
        if app:
            conn.execute("UPDATE member_applications SET status = 'rejected' WHERE id = ?", (app_id,))
            conn.execute("DELETE FROM member_application_pending WHERE user_id = ?", (app["user_id"],))
            u = conn.execute("SELECT tg_user_id FROM users WHERE id = ?", (app["user_id"],)).fetchone()
        else:
            u = None
    if u:
        u = dict(u)
    if u and u.get("tg_user_id"):
        await _send_system_notice(
            int(u["tg_user_id"]),
            "Вашу заявку на участь у спільноті відхилено. Якщо це помилка — напишіть у «Питання адміну».",
        )
    return RedirectResponse(url=f"/admin/applications?token={token}", status_code=303)


@app.post("/admin/users/{user_id}/role")
async def admin_change_role(user_id: int, token: str = Form(...), role: str = Form(...)):
    _verify_admin_token(token)
    if role in ("guest", "member", "lead"):
        with db() as conn:
            conn.execute("UPDATE users SET role = ? WHERE id = ?", (role, user_id))
            u = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if u:
            u = dict(u)
            await sendpulse_upsert({"external_id": str(u["id"]), "full_name": u.get("full_name") or "", "phone": u.get("phone") or "", "email": u.get("email") or "", "role": role, "telegram_id": u.get("tg_user_id"), "tg_username": u.get("tg_username") or "", "payment_status": "", "event_ids": ""})
    return RedirectResponse(url=f"/admin?token={token}", status_code=303)


@app.get("/admin/membership", response_class=HTMLResponse)
async def admin_membership_page(request: Request, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        membership_settings = conn.execute("SELECT * FROM membership_settings ORDER BY months").fetchall()
    wayforpay_membership_url = _get_app_config("wayforpay_membership_url", "")
    return templates.TemplateResponse(
        "admin_membership.html",
        {"request": request, "token": token, "settings": [dict(s) for s in membership_settings], "wayforpay_membership_url": wayforpay_membership_url},
    )


@app.post("/admin/membership/update")
async def admin_membership_update(token: str = Form(...), period: str = Form(...), amount_uah: int = Form(...)):
    _verify_admin_token(token)
    with db() as conn:
        conn.execute("UPDATE membership_settings SET amount_uah = ? WHERE period = ?", (amount_uah, period))
    return RedirectResponse(url=f"/admin/membership?token={token}", status_code=303)


@app.post("/admin/membership/payment-url")
async def admin_membership_payment_url(token: str = Form(...), wayforpay_membership_url: str = Form("")):
    _verify_admin_token(token)
    _set_app_config("wayforpay_membership_url", wayforpay_membership_url.strip())
    return RedirectResponse(url=f"/admin/membership?token={token}", status_code=303)


@app.get("/admin/bot-menu", response_class=HTMLResponse)
async def admin_bot_menu_page(request: Request, token: str = ""):
    _verify_admin_token(token)
    ctx = {
        "request": request,
        "token": token,
        "academy_about_text": _get_app_config("academy_about_text", ""),
        "policy_url": _get_app_config("policy_url", ""),
        "manager_contact": _get_app_config("manager_contact", ""),
        "telegram_bot_username": _get_app_config("telegram_bot_username", ""),
        "default_msg_rsvp_40h": _get_app_config("default_msg_rsvp_40h", ""),
        "default_msg_remind_24h": _get_app_config("default_msg_remind_24h", ""),
        "default_msg_remind_3h": _get_app_config("default_msg_remind_3h", ""),
        "timeline_rsvp_40_enabled": _get_app_config("timeline_rsvp_40_enabled", "1"),
        "timeline_24h_enabled": _get_app_config("timeline_24h_enabled", "1"),
        "timeline_3h_enabled": _get_app_config("timeline_3h_enabled", "1"),
        "quiet_hours_enabled": _get_app_config("quiet_hours_enabled", "0"),
        "quiet_hours_start": _get_app_config("quiet_hours_start", "22"),
        "quiet_hours_end": _get_app_config("quiet_hours_end", "8"),
        "quiet_hours_tz": _get_app_config("quiet_hours_tz", "Europe/Kyiv"),
    }
    return templates.TemplateResponse("admin_bot_menu.html", ctx)


@app.post("/admin/bot-menu/update")
async def admin_bot_menu_update(
    token: str = Form(...),
    academy_about_text: str = Form(""),
    policy_url: str = Form(""),
    manager_contact: str = Form(""),
    telegram_bot_username: str = Form(""),
    default_msg_rsvp_40h: str = Form(""),
    default_msg_remind_24h: str = Form(""),
    default_msg_remind_3h: str = Form(""),
    timeline_rsvp_40_enabled: str = Form("1"),
    timeline_24h_enabled: str = Form("1"),
    timeline_3h_enabled: str = Form("1"),
    quiet_hours_enabled: str = Form("0"),
    quiet_hours_start: str = Form("22"),
    quiet_hours_end: str = Form("8"),
    quiet_hours_tz: str = Form("Europe/Kyiv"),
):
    _verify_admin_token(token)
    _set_app_config("academy_about_text", academy_about_text)
    _set_app_config("policy_url", policy_url.strip())
    _set_app_config("manager_contact", manager_contact.strip())
    _set_app_config("telegram_bot_username", telegram_bot_username.strip().lstrip("@"))
    _set_app_config("default_msg_rsvp_40h", default_msg_rsvp_40h.strip())
    _set_app_config("default_msg_remind_24h", default_msg_remind_24h.strip())
    _set_app_config("default_msg_remind_3h", default_msg_remind_3h.strip())
    _set_app_config("timeline_rsvp_40_enabled", (timeline_rsvp_40_enabled or "1").strip()[:1])
    _set_app_config("timeline_24h_enabled", (timeline_24h_enabled or "1").strip()[:1])
    _set_app_config("timeline_3h_enabled", (timeline_3h_enabled or "1").strip()[:1])
    _set_app_config("quiet_hours_enabled", (quiet_hours_enabled or "0").strip()[:1])
    _set_app_config("quiet_hours_start", (quiet_hours_start or "22").strip()[:2])
    _set_app_config("quiet_hours_end", (quiet_hours_end or "8").strip()[:2])
    _set_app_config("quiet_hours_tz", (quiet_hours_tz or "Europe/Kyiv").strip()[:80])
    return RedirectResponse(url=f"/admin/bot-menu?token={token}", status_code=303)


@app.get("/admin/tags", response_class=HTMLResponse)
async def admin_tags(request: Request, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        tags = conn.execute("SELECT * FROM member_tags ORDER BY name").fetchall()
    return templates.TemplateResponse("admin_tags.html", {"request": request, "token": token, "tags": [dict(t) for t in tags]})


@app.post("/admin/tags/create")
async def admin_create_tag(token: str = Form(...), name: str = Form(...)):
    _verify_admin_token(token)
    with db() as conn:
        conn.execute("INSERT OR IGNORE INTO member_tags (name) VALUES (?)", (name.strip(),))
    return RedirectResponse(url=f"/admin/tags?token={token}", status_code=303)


@app.get("/admin/push", response_class=HTMLResponse)
async def admin_push_page(request: Request, token: str = "", user_id: str = "", event_id: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id = ?", (int(user_id),)).fetchone() if user_id else None
        events = conn.execute("SELECT id, title FROM events WHERE status='upcoming' ORDER BY starts_at").fetchall() if event_id else []
    user_dict = dict(user) if user else None
    if user_dict:
        user_dict["tg_display"] = format_tg_display(user_dict)
    return templates.TemplateResponse(
        "admin_push.html",
        {"request": request, "token": token, "user": user_dict, "event_id": event_id, "events": [dict(e) for e in events]},
    )


@app.post("/admin/push/send")
async def admin_push_send(
    token: str = Form(...),
    user_id: int = Form(...),
    message: str = Form(""),
    event_id: str = Form(""),
    media_type: str = Form("text"),
    attachment_file: UploadFile = File(None),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_admin_token(token)
    text = (message or "").strip()
    uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
    if not uploads and attachment_file and attachment_file.filename:
        uploads = [attachment_file]
    with db() as conn:
        u = conn.execute("SELECT tg_user_id FROM users WHERE id = ?", (user_id,)).fetchone()
    if u:
        if uploads:
            first = True
            for up in uploads:
                media_url, content_type = await _save_uploaded_media(up, f"push_{user_id}")
                if not media_url:
                    continue
                cap = text[:1024] if first else ""
                if content_type == "photo":
                    await send_photo(u["tg_user_id"], media_url, caption=cap or None)
                elif content_type == "video":
                    await send_video(u["tg_user_id"], media_url, caption=cap or None)
                else:
                    await send_document(u["tg_user_id"], media_url, caption=cap or None)
                first = False
        elif text:
            await send_message(u["tg_user_id"], text[:4096])
    if event_id:
        return RedirectResponse(url=f"/admin/events/{event_id}?token={token}", status_code=303)
    return RedirectResponse(url=f"/admin?token={token}", status_code=303)


@app.get("/admin/support", response_class=HTMLResponse)
async def admin_support_page(
    request: Request,
    token: str = "",
    st: str = "",
    thread_id: str = "",
    tab: str = "active",
):
    st_s, tok_s = (st or "").strip(), (token or "").strip()
    st_d = _decode_signed_token(st_s) if st_s else None
    ad = _decode_signed_token(tok_s) if tok_s else None
    full_admin = bool(ad and "tg_user_id" in ad)
    st_ok = bool(st_d and st_d.get("purpose") == "support_thread")
    if not full_admin and not st_ok:
        return templates.TemplateResponse(
            "admin_login.html",
            {"request": request, "error": "Невалідне або прострочене посилання. Відкрийте /admin у боті."},
        )
    effective_token = tok_s if full_admin else st_s
    tid_st = int(st_d["thread_id"]) if st_ok else None
    if st_ok and tid_st is not None:
        if thread_id and str(int(thread_id)) != str(tid_st):
            return templates.TemplateResponse(
                "admin_login.html",
                {"request": request, "error": "Посилання не відповідає цьому чату."},
            )
        thread_id = str(tid_st)
    if full_admin:
        with db() as conn:
            orphans = conn.execute(
                """SELECT u.id FROM users u WHERE u.onboarding_stage = 'waiting_support'
                AND NOT EXISTS (SELECT 1 FROM support_threads st WHERE st.user_id = u.id AND st.status = 'open')"""
            ).fetchall()
        for row in orphans:
            ensure_open_support_thread(int(row["id"]))
    threads_data: list = []
    archived_threads_data: list = []
    if full_admin:
        with db() as conn:
            threads = conn.execute(
                """SELECT st.id, st.user_id, st.status, st.updated_at, u.full_name, u.tg_user_id, u.tg_username
                FROM support_threads st JOIN users u ON u.id = st.user_id
                WHERE st.status = 'open' ORDER BY st.updated_at DESC"""
            ).fetchall()
            archived = conn.execute(
                """SELECT st.id, st.user_id, st.status, st.updated_at, u.full_name, u.tg_user_id, u.tg_username
                FROM support_threads st JOIN users u ON u.id = st.user_id
                WHERE st.status = 'archived' ORDER BY st.updated_at DESC LIMIT 200"""
            ).fetchall()
        for t in threads:
            d = dict(t)
            d["tg_display"] = format_tg_display(d)
            threads_data.append(d)
        for t in archived:
            d = dict(t)
            d["tg_display"] = format_tg_display(d)
            archived_threads_data.append(d)
    thread = None
    user: dict = {}
    messages: list = []
    thread_archived = False
    if thread_id:
        try:
            tid = int(thread_id)
        except ValueError:
            tid = 0
        if tid:
            _verify_support_route_token(effective_token, tid)
            with db() as conn:
                thread = conn.execute(
                    "SELECT * FROM support_threads WHERE id = ? AND status IN ('open', 'archived')",
                    (tid,),
                ).fetchone()
            if thread:
                thread = dict(thread)
                thread_archived = thread.get("status") == "archived"
                with db() as conn:
                    urow = conn.execute("SELECT * FROM users WHERE id = ?", (thread["user_id"],)).fetchone()
                user = dict(urow) if urow else {}
                user["tg_display"] = format_tg_display(user)
                with db() as conn:
                    mrows = conn.execute(
                        "SELECT * FROM support_messages WHERE thread_id = ? ORDER BY created_at",
                        (thread["id"],),
                    ).fetchall()
                messages = [dict(m) for m in mrows]
    return templates.TemplateResponse(
        "admin_support.html",
        {
            "request": request,
            "token": effective_token,
            "full_admin": full_admin,
            "threads": threads_data,
            "archived_threads": archived_threads_data,
            "thread": thread,
            "user": user,
            "messages": messages,
            "tab": (tab or "active").strip() or "active",
            "thread_archived": thread_archived,
        },
    )


@app.post("/admin/support/{thread_id}/reply")
async def admin_support_reply(
    thread_id: int,
    token: str = Form(...),
    message: str = Form(""),
    media_type: str = Form("text"),
    photo_file: UploadFile = File(None),
    video_file: UploadFile = File(None),
    attachment_file: UploadFile = File(None),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_support_route_token(token, thread_id)
    with db() as conn:
        thread = conn.execute("SELECT * FROM support_threads WHERE id = ? AND status = 'open'", (thread_id,)).fetchone()
        if not thread:
            raise HTTPException(status_code=404, detail="Thread not found")
        user = conn.execute("SELECT tg_user_id FROM users WHERE id = ?", (thread["user_id"],)).fetchone()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    now = utcnow_iso()
    caption = (message or "").strip()[:1024]
    uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
    if not uploads and attachment_file and attachment_file.filename:
        uploads = [attachment_file]
    if not uploads:
        if media_type == "photo" and photo_file and photo_file.filename:
            uploads = [photo_file]
        elif media_type == "video" and video_file and video_file.filename:
            uploads = [video_file]

    sent_any = False
    if uploads:
        first = True
        for up in uploads:
            media_url, content_type = await _save_uploaded_media(up, f"sup_{thread_id}")
            if not media_url:
                continue
            cap = caption if first else ""
            if content_type == "photo":
                await send_photo(user["tg_user_id"], media_url, caption=cap or None)
            elif content_type == "video":
                await send_video(user["tg_user_id"], media_url, caption=cap or None)
            else:
                await send_document(user["tg_user_id"], media_url, caption=cap or None)
            with db() as conn:
                conn.execute(
                    "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, ?, ?)",
                    (thread_id, media_url, content_type, now),
                )
            sent_any = True
            first = False
    elif caption:
        await send_message(user["tg_user_id"], caption[:4096])
        with db() as conn:
            conn.execute(
                "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, 'text', ?)",
                (thread_id, caption[:4096], now),
            )
        sent_any = True

    if sent_any:
        with db() as conn:
            conn.execute("UPDATE support_threads SET updated_at = ? WHERE id = ?", (now, thread_id))
    q = urlencode({"token": token, "thread_id": str(thread_id)})
    d = _decode_signed_token(token)
    if d and d.get("purpose") == "support_thread":
        q = urlencode({"st": token, "thread_id": str(thread_id)})
    return RedirectResponse(url=f"/admin/support?{q}", status_code=303)


async def _admin_support_archive_thread(thread_id: int, token: str) -> None:
    _verify_support_route_token(token, thread_id)
    tg_uid: int | None = None
    now = utcnow_iso()
    with db() as conn:
        row = conn.execute(
            "SELECT user_id FROM support_threads WHERE id = ? AND status = 'open'",
            (thread_id,),
        ).fetchone()
        if row:
            conn.execute(
                "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, 'bot_outbound', ?)",
                (thread_id, SUPPORT_USER_GOODBYE_TEXT, now),
            )
            conn.execute("UPDATE support_threads SET status = 'archived', updated_at = ? WHERE id = ?", (now, thread_id))
            conn.execute("UPDATE users SET onboarding_stage = 'done' WHERE id = ?", (row["user_id"],))
            u = conn.execute("SELECT tg_user_id FROM users WHERE id = ?", (row["user_id"],)).fetchone()
            if u:
                tg_uid = int(u["tg_user_id"])
    if tg_uid:
        await send_message(tg_uid, SUPPORT_USER_GOODBYE_TEXT)


@app.post("/admin/support/{thread_id}/archive")
async def admin_support_archive(thread_id: int, token: str = Form(...)):
    await _admin_support_archive_thread(thread_id, token)
    d = _decode_signed_token(token)
    if d and d.get("purpose") == "support_thread":
        return RedirectResponse(url=f"/admin/support?{urlencode({'st': token, 'thread_id': str(thread_id)})}", status_code=303)
    return RedirectResponse(url=f"/admin/support?{urlencode({'token': token, 'tab': 'archive'})}", status_code=303)


@app.post("/admin/support/{thread_id}/close")
async def admin_support_close(thread_id: int, token: str = Form(...)):
    await _admin_support_archive_thread(thread_id, token)
    d = _decode_signed_token(token)
    if d and d.get("purpose") == "support_thread":
        return RedirectResponse(url=f"/admin/support?{urlencode({'st': token, 'thread_id': str(thread_id)})}", status_code=303)
    return RedirectResponse(url=f"/admin/support?{urlencode({'token': token, 'tab': 'archive'})}", status_code=303)


@app.get("/admin/support/{thread_id}/messages")
async def admin_support_messages(thread_id: int, token: str = "", after: int = 0):
    _verify_support_route_token(token, thread_id)
    with db() as conn:
        rows = conn.execute(
            "SELECT id, from_admin, content, content_type, created_at FROM support_messages WHERE thread_id = ? AND id > ? ORDER BY id",
            (thread_id, after),
        ).fetchall()
    return JSONResponse({
        "messages": [
            {
                "id": r["id"],
                "kind": _support_message_kind(r),
                "from_admin": bool(r["from_admin"]),
                "content": r["content"],
                "content_type": r["content_type"],
                "created_at": _format_datetime(r["created_at"]),
            }
            for r in rows
        ]
    })


def _role_label_uk(role: str | None) -> str:
    return {"member": "Учасник", "guest": "Гість", "lead": "Лід"}.get((role or "").strip(), (role or "").strip() or "—")


def _support_message_kind(row) -> str:
    ct = (row["content_type"] if hasattr(row, "keys") else row.get("content_type")) or ""
    ct = str(ct).strip()
    if ct == "bot_outbound":
        return "bot"
    fa = row["from_admin"] if hasattr(row, "keys") else row.get("from_admin")
    return "admin" if fa else "user"


@app.get("/admin/live-chats", response_class=HTMLResponse)
async def admin_live_chats_page(request: Request, token: str = "", user_id: str = ""):
    _verify_admin_token(token)
    sel: int | None = None
    s = (user_id or "").strip()
    if s.isdigit():
        sel = int(s)
    return templates.TemplateResponse(
        "admin_live_chats.html",
        {"request": request, "token": token, "selected_user_id": sel},
    )


@app.get("/admin/live-chats/data")
async def admin_live_chats_data(token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        rows = conn.execute(
            """
            SELECT u.id, u.full_name, u.phone, u.role, u.onboarding_stage, u.updated_at, u.tg_user_id, u.tg_username,
              (SELECT GROUP_CONCAT(e.title || ' [' || COALESCE(r.status, '') || ']', char(10))
                 FROM registrations r
                 JOIN events e ON e.id = r.event_id AND e.status = 'upcoming'
                WHERE r.user_id = u.id AND r.status IN ('registered', 'pending_payment')
              ) AS reg_events,
              (SELECT st.id FROM support_threads st WHERE st.user_id = u.id AND st.status = 'open' ORDER BY st.id DESC LIMIT 1) AS open_thread_id,
              (SELECT sm.content FROM support_messages sm
                 JOIN support_threads st2 ON sm.thread_id = st2.id
                WHERE st2.user_id = u.id AND st2.status = 'open'
                ORDER BY sm.id DESC LIMIT 1) AS last_snippet,
              (SELECT sm.created_at FROM support_messages sm
                 JOIN support_threads st3 ON sm.thread_id = st3.id
                WHERE st3.user_id = u.id AND st3.status = 'open'
                ORDER BY sm.id DESC LIMIT 1) AS last_msg_at
            FROM users u
            ORDER BY
              CASE WHEN (SELECT st4.id FROM support_threads st4 WHERE st4.user_id = u.id AND st4.status = 'open' ORDER BY st4.id DESC LIMIT 1) IS NOT NULL THEN 0 ELSE 1 END,
              COALESCE(
                (SELECT sm2.created_at FROM support_messages sm2
                   JOIN support_threads st5 ON sm2.thread_id = st5.id
                  WHERE st5.user_id = u.id AND st5.status = 'open'
                  ORDER BY sm2.id DESC LIMIT 1),
                u.updated_at, ''
              ) DESC,
              u.id DESC
            LIMIT 400
            """
        ).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        d["tg_display"] = format_tg_display(d)
        d["role_ua"] = _role_label_uk(d.get("role"))
        ev = (d.get("reg_events") or "").strip()
        d["events_lines"] = [ln.strip() for ln in ev.split("\n") if ln.strip()] if ev else []
        d.pop("reg_events", None)
        tid = d.get("open_thread_id")
        d["open_thread_id"] = int(tid) if tid is not None else None
        out.append(d)
    return JSONResponse({"users": out})


@app.get("/admin/live-chats/{user_id}/messages")
async def admin_live_chats_user_messages(user_id: int, token: str = "", after: int = 0):
    _verify_admin_token(token)
    with db() as conn:
        row = conn.execute(
            "SELECT id FROM support_threads WHERE user_id = ? AND status = 'open' ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
        if not row:
            return JSONResponse({"messages": [], "thread_id": None})
        tid = int(row["id"])
        mrows = conn.execute(
            "SELECT id, from_admin, content, content_type, created_at FROM support_messages WHERE thread_id = ? AND id > ? ORDER BY id",
            (tid, after),
        ).fetchall()
    return JSONResponse(
        {
            "thread_id": tid,
            "messages": [
                {
                    "id": r["id"],
                    "kind": _support_message_kind(r),
                    "from_admin": bool(r["from_admin"]),
                    "content": r["content"],
                    "content_type": r["content_type"],
                    "created_at": _format_datetime(r["created_at"]),
                }
                for r in mrows
            ],
        }
    )


@app.post("/admin/live-chats/{user_id}/send")
async def admin_live_chats_send(
    user_id: int,
    token: str = Form(...),
    message: str = Form(""),
    attachment_file: UploadFile = File(None),
    media_type: str = Form("text"),
    attachment_files: list[UploadFile] = File(None),
):
    _verify_admin_token(token)
    text = (message or "").strip()
    q = urlencode({"token": token, "user_id": str(user_id)})
    uploads: list[UploadFile] = [f for f in (attachment_files or []) if f and f.filename]
    if not uploads and attachment_file and attachment_file.filename:
        uploads = [attachment_file]
    has_attachment = bool(uploads)
    if not text and not has_attachment:
        return RedirectResponse(url=f"/admin/live-chats?{q}", status_code=303)
    now = utcnow_iso()
    tg_target: int | None = None
    with db() as conn:
        u = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if not u:
            raise HTTPException(status_code=404, detail="User not found")
        tg_target = int(u["tg_user_id"])
        row = conn.execute(
            "SELECT id FROM support_threads WHERE user_id = ? AND status = 'open' ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
        if row:
            tid = int(row["id"])
            conn.execute("UPDATE support_threads SET updated_at = ? WHERE id = ?", (now, tid))
        else:
            conn.execute(
                "INSERT INTO support_threads (user_id, status, created_at, updated_at) VALUES (?, 'open', ?, ?)",
                (user_id, now, now),
            )
            tid = int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])
    sent_any = False
    if uploads:
        first = True
        for up in uploads:
            media_url, content_type = await _save_uploaded_media(up, f"livechat_{user_id}")
            if not media_url:
                continue
            cap = text[:1024] if first else ""
            if content_type == "photo":
                await send_photo(tg_target, media_url, caption=cap or None)
            elif content_type == "video":
                await send_video(tg_target, media_url, caption=cap or None)
            else:
                await send_document(tg_target, media_url, caption=cap or None)
            with db() as conn:
                conn.execute(
                    "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, ?, ?)",
                    (tid, media_url, content_type, now),
                )
            sent_any = True
            first = False
    elif text:
        await send_message(tg_target, text[:4096])
        with db() as conn:
            conn.execute(
                "INSERT INTO support_messages (thread_id, from_admin, content, content_type, created_at) VALUES (?, 1, ?, 'text', ?)",
                (tid, text[:4096], now),
            )
        sent_any = True

    if sent_any:
        with db() as conn:
            conn.execute("UPDATE support_threads SET updated_at = ? WHERE id = ?", (now, tid))
    return RedirectResponse(url=f"/admin/live-chats?{q}", status_code=303)


@app.get("/admin/reminders", response_class=HTMLResponse)
async def admin_reminders_page(request: Request, token: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        settings_rows = conn.execute(
            "SELECT * FROM reminder_settings ORDER BY stage"
        ).fetchall()
        events = conn.execute("SELECT id, title, starts_at FROM events WHERE status='upcoming' ORDER BY starts_at").fetchall()
    return templates.TemplateResponse(
        "admin_reminders.html",
        {"request": request, "token": token, "settings": [dict(s) for s in settings_rows], "events": [dict(e) for e in events]},
    )


@app.post("/admin/reminders/update")
async def admin_reminders_update(
    token: str = Form(...),
    stage: str = Form(...),
    message_text: str = Form(...),
    delay_minutes: int = Form(60),
    max_attempts: int = Form(3),
    enabled: str = Form("1"),
):
    _verify_admin_token(token)
    with db() as conn:
        conn.execute(
            """UPDATE reminder_settings SET message_text=?, delay_minutes=?, max_attempts=?, enabled=?
            WHERE stage=?""",
            (message_text, delay_minutes, max_attempts, 1 if enabled == "1" else 0, stage),
        )
    return RedirectResponse(url=f"/admin/reminders?token={token}", status_code=303)


@app.get("/admin/funnel", response_class=HTMLResponse)
async def admin_funnel_page(request: Request, token: str = "", stage: str = ""):
    _verify_admin_token(token)
    with db() as conn:
        # Users by funnel stage
        stage_1 = conn.execute(
            """SELECT u.id, u.full_name, u.tg_username, u.tg_user_id, u.onboarding_stage, u.updated_at
            FROM users u WHERE u.onboarding_stage NOT IN ('done', 'new') AND u.onboarding_stage != ''
            ORDER BY u.updated_at DESC LIMIT 100"""
        ).fetchall()
        stage_2 = conn.execute(
            """SELECT u.id, u.full_name, u.tg_username, u.tg_user_id, u.updated_at
            FROM users u WHERE u.onboarding_stage = 'done'
            AND NOT EXISTS (SELECT 1 FROM registrations r WHERE r.user_id = u.id)
            ORDER BY u.updated_at DESC LIMIT 100"""
        ).fetchall()
        stage_4 = conn.execute(
            """SELECT u.id, u.full_name, u.tg_username, u.tg_user_id, r.event_id, r.wayforpay_order_reference, e.title as event_title
            FROM registrations r JOIN users u ON u.id = r.user_id
            JOIN events e ON e.id = r.event_id
            WHERE r.payment_status NOT IN ('paid', 'free') AND r.wayforpay_order_reference IS NOT NULL
            AND e.status = 'upcoming' ORDER BY r.updated_at DESC LIMIT 100"""
        ).fetchall()
        stage_5 = conn.execute(
            """SELECT u.id, u.full_name, u.tg_username, u.tg_user_id, r.event_id, e.title as event_title, e.starts_at
            FROM registrations r JOIN users u ON u.id = r.user_id
            JOIN events e ON e.id = r.event_id
            WHERE r.payment_status IN ('paid', 'free') AND r.status = 'registered'
            AND e.status = 'upcoming' ORDER BY e.starts_at LIMIT 100"""
        ).fetchall()
    return templates.TemplateResponse(
        "admin_funnel.html",
        {
            "request": request,
            "token": token,
            "stage_1": [dict(r) for r in stage_1],
            "stage_2": [dict(r) for r in stage_2],
            "stage_4": [dict(r) for r in stage_4],
            "stage_5": [dict(r) for r in stage_5],
        },
    )


@app.get("/admin/import", response_class=HTMLResponse)
async def admin_import_page(request: Request, token: str = "", imported: str = ""):
    _verify_admin_token(token)
    return templates.TemplateResponse("admin_import.html", {"request": request, "token": token, "imported": imported or None})


@app.post("/admin/import/excel")
async def admin_import_excel(token: str = Form(...), file: UploadFile = Form(...)):
    _verify_admin_token(token)
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Завантажте файл .xlsx з колонкою 'phone' або 'телефон'")
    buf = await file.read()
    wb = openpyxl.load_workbook(filename=io.BytesIO(buf), read_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    phone_col = None
    for i, h in enumerate(headers):
        if h in ("phone", "телефон", "phones", "номер"):
            phone_col = i
            break
    if phone_col is None:
        wb.close()
        raise HTTPException(status_code=400, detail="Не знайдено колонку phone/телефон")
    def _digits_from_cell(v) -> str:
        if v is None:
            return ""
        if isinstance(v, int):
            return str(v)
        if isinstance(v, float):
            # Excel often stores phone as number; avoid scientific notation and trailing .0.
            return f"{v:.0f}"
        return str(v)

    def _normalize_ua_phone(raw: str) -> str:
        d = "".join(ch for ch in (raw or "") if ch.isdigit())
        if not d:
            return ""
        # +380XXXXXXXXX / 380XXXXXXXXX
        if d.startswith("380") and len(d) == 12:
            return d
        # 0XXXXXXXXX -> 380XXXXXXXXX
        if d.startswith("0") and len(d) == 10:
            return "380" + d[1:]
        # XXXXXXXXX (without country code)
        if len(d) == 9:
            return "380" + d
        # Fallback for messy values: try by suffix.
        if len(d) > 12 and d.endswith("380") is False:
            tail = d[-9:]
            if len(tail) == 9:
                return "380" + tail
        if len(d) >= 12 and d.startswith("380"):
            return d[:12]
        return ""

    phones: set[str] = set()
    for row in ws.iter_rows(min_row=2):
        val = row[phone_col].value
        if val:
            normalized = _normalize_ua_phone(_digits_from_cell(val))
            if normalized:
                phones.add(normalized)
    wb.close()

    updated = 0
    with db() as conn:
        all_users = conn.execute("SELECT id, phone FROM users").fetchall()
        for u in all_users:
            un = _normalize_ua_phone(str(u["phone"] or ""))
            if not un or un not in phones:
                continue
            conn.execute("UPDATE users SET role = 'member' WHERE id = ?", (u["id"],))
            updated += 1
    return RedirectResponse(url=f"/admin/import?token={token}&imported={updated}", status_code=303)


@app.get("/admin/export")
async def admin_export(token: str = "", segment: str = "all", event_id: str = ""):
    _verify_admin_token(token)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    with db() as conn:
        if segment == "event" and event_id:
            rows = conn.execute(
                """SELECT u.id, u.full_name, u.phone, u.email, u.role, u.tg_username, u.tg_user_id, r.payment_status
                FROM registrations r JOIN users u ON u.id = r.user_id WHERE r.event_id = ?""",
                (int(event_id),),
            ).fetchall()
            headers = ["id", "full_name", "phone", "email", "role", "tg_username", "tg_user_id", "payment_status"]
        elif segment == "members":
            rows = conn.execute("SELECT id, full_name, phone, email, role, tg_username, tg_user_id FROM users WHERE role = 'member'").fetchall()
            headers = ["id", "full_name", "phone", "email", "role", "tg_username", "tg_user_id"]
        elif segment == "guests":
            rows = conn.execute("SELECT id, full_name, phone, email, role, tg_username, tg_user_id FROM users WHERE role = 'guest'").fetchall()
            headers = ["id", "full_name", "phone", "email", "role", "tg_username", "tg_user_id"]
        else:
            rows = conn.execute("SELECT id, full_name, phone, email, role, tg_username, tg_user_id FROM users").fetchall()
            headers = ["id", "full_name", "phone", "email", "role", "tg_username", "tg_user_id"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    for row_idx, r in enumerate(rows, 2):
        for col_idx, k in enumerate(headers, 1):
            val = r[k] if k in r.keys() else ""
            ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=export.xlsx"},
    )


async def _send_admin_event_digest(event_id: int, evd: dict) -> None:
    """Один раз за ~24 год до старту: зведення для адмінів Telegram."""
    seat_sql = _sql_registration_occupies_seat()
    title = evd.get("title") or "Подія"
    with db() as conn:
        chk = conn.execute("SELECT admin_digest_24h_sent FROM events WHERE id = ?", (event_id,)).fetchone()
        if chk and int(chk["admin_digest_24h_sent"] or 0):
            return
        occ = conn.execute(
            f"SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND {seat_sql}",
            (event_id,),
        ).fetchone()["c"]
        paid_confirmed = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND status = 'registered'
            AND payment_status IN ('paid','free') AND IFNULL(rsvp_will_attend,0) = 1""",
            (event_id,),
        ).fetchone()["c"]
        paid_unconf = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ? AND status = 'registered'
            AND payment_status IN ('paid','free') AND rsvp_will_attend IS NULL""",
            (event_id,),
        ).fetchone()["c"]
        unpaid = conn.execute(
            """SELECT COUNT(*) as c FROM registrations WHERE event_id = ?
            AND NOT (status = 'registered' AND payment_status IN ('paid','free'))""",
            (event_id,),
        ).fetchone()["c"]
    lim = int(evd.get("seats_limit") or 0)
    free_seats = max(0, lim - occ) if lim else None
    lines = [
        f"📊 Зведення за ~24 год до події «{title}»",
        f"✅ Оплатили й підтвердили «Буду»: {paid_confirmed}",
        f"⚠️ Оплатили, без підтвердження участі: {paid_unconf}",
        f"💳 Не оплатили / не завершено: {unpaid}",
    ]
    if free_seats is not None:
        lines.append(f"🪑 Вільних місць (за лімітом): {free_seats}")
    msg = "\n".join(lines)
    for aid in settings.admin_ids:
        try:
            await send_message(aid, msg)
        except Exception:
            pass
    with db() as conn:
        conn.execute("UPDATE events SET admin_digest_24h_sent = 1 WHERE id = ?", (event_id,))


async def event_timeline_worker():
    """Нагадування за ТЗ: ~40 год RSVP, ~24 год, ~3 год; шаблони; тихі години; перемикачі в app_config."""
    while True:
        try:
            now = datetime.now(timezone.utc)
            quiet = _is_quiet_hours_now()
            with db() as conn:
                events = conn.execute("SELECT * FROM events WHERE status = 'upcoming'").fetchall()
            for ev in events:
                evd = dict(ev)
                try:
                    starts = datetime.fromisoformat(str(evd.get("starts_at") or "").replace("Z", "+00:00"))
                except Exception:
                    continue
                if starts.tzinfo is None:
                    starts = starts.replace(tzinfo=timezone.utc)
                delta = starts - now
                if delta.total_seconds() <= 0:
                    continue
                eid = int(evd["id"])

                if timedelta(hours=22) <= delta <= timedelta(hours=26) and not int(evd.get("admin_digest_24h_sent") or 0):
                    await _send_admin_event_digest(eid, evd)

                with db() as conn:
                    regs = conn.execute(
                        """SELECT r.*, u.tg_user_id FROM registrations r
                        JOIN users u ON u.id = r.user_id
                        WHERE r.event_id = ? AND r.status = 'registered' AND r.payment_status IN ('paid','free')""",
                        (eid,),
                    ).fetchall()
                for r in regs:
                    rd = dict(r)
                    tg = rd.get("tg_user_id")
                    if not tg:
                        continue
                    rid = int(rd["id"])
                    title = evd.get("title") or "Подія"
                    dt_human = _format_datetime(str(evd.get("starts_at") or ""))
                    loc = (evd.get("location") or "").strip()
                    zoom = (evd.get("zoom_url") or "").strip()

                    if int(rd.get("reminder_rsvp_40h_sent") or 0) == 0 and timedelta(hours=36) <= delta <= timedelta(hours=44) and rd.get("rsvp_will_attend") is None:
                        if not _timeline_send_allowed("rsvp_40"):
                            continue
                        if quiet:
                            continue
                        m = _format_event_reminder_body(evd, "rsvp_40", title, dt_human, loc, zoom)
                        await send_message(
                            int(tg),
                            m,
                            reply_markup=inline_keyboard(
                                [
                                    [
                                        {"text": "✅ Буду", "callback_data": f"rsvp_yes_{eid}"},
                                        {"text": "❌ Не буду", "callback_data": f"rsvp_no_{eid}"},
                                    ]
                                ]
                            ),
                        )
                        with db() as conn:
                            conn.execute(
                                "UPDATE registrations SET reminder_rsvp_40h_sent = 1, updated_at = ? WHERE id = ?",
                                (utcnow_iso(), rid),
                            )
                        continue

                    attending = rd.get("rsvp_will_attend") is None or int(rd.get("rsvp_will_attend") or 0) == 1
                    if not attending:
                        continue

                    if int(rd.get("reminder_24h_sent") or 0) == 0 and timedelta(hours=22) <= delta <= timedelta(hours=26):
                        if not _timeline_send_allowed("24h"):
                            continue
                        if quiet:
                            continue
                        body = _format_event_reminder_body(evd, "24h", title, dt_human, loc, zoom)
                        await send_message(int(tg), body)
                        with db() as conn:
                            conn.execute(
                                "UPDATE registrations SET reminder_24h_sent = 1, updated_at = ? WHERE id = ?",
                                (utcnow_iso(), rid),
                            )

                    if int(rd.get("reminder_3h_sent") or 0) == 0 and timedelta(hours=2, minutes=30) <= delta <= timedelta(hours=3, minutes=30):
                        if not _timeline_send_allowed("3h"):
                            continue
                        if quiet:
                            continue
                        short = _format_event_reminder_body(evd, "3h", title, dt_human, loc, zoom)
                        if loc.startswith("http"):
                            await send_message(
                                int(tg),
                                short,
                                reply_markup=inline_keyboard([[{"text": "🗺 Локація / карта", "url": loc}]]),
                            )
                        elif zoom:
                            await send_message(int(tg), f"{short}\n\n🔗 {zoom}")
                        else:
                            await send_message(int(tg), short + (f"\n📍 {loc}" if loc else ""))
                        with db() as conn:
                            conn.execute(
                                "UPDATE registrations SET reminder_3h_sent = 1, updated_at = ? WHERE id = ?",
                                (utcnow_iso(), rid),
                            )
        except Exception:
            pass
        await asyncio.sleep(120)


def _should_skip_reminder(r: dict) -> bool:
    """Перевіряє, чи потрібно пропустити нагадування (користувач вже перейшов на наступний етап)."""
    with db() as conn:
        if r["stage"] == "stage_1":
            u = conn.execute("SELECT onboarding_stage FROM users WHERE id = ?", (r["user_id"],)).fetchone()
            return u and str(u["onboarding_stage"] or "") == "done"
        if r["stage"] == "stage_2":
            has_reg = conn.execute("SELECT 1 FROM registrations WHERE user_id = ? LIMIT 1", (r["user_id"],)).fetchone()
            return bool(has_reg)
        if r["stage"] in ("stage_4_unpaid_24h", "stage_4_unpaid_1h") and r.get("event_id"):
            reg = get_registration(r["user_id"], r["event_id"])
            return bool(reg and reg["payment_status"] in ("paid", "free"))
    return False


async def reminder_worker():
    while True:
        try:
            rows = due_reminders()
            for r in rows:
                if _should_skip_reminder(dict(r)):
                    mark_sent(r["id"])
                    continue
                if r["stage"] in ("stage_4", "stage_4_unpaid_24h", "stage_4_unpaid_1h") and r.get("event_id"):
                    reg = get_registration(r["user_id"], r["event_id"])
                    reg = row_to_dict(reg) if reg else None
                    if reg and reg.get("payment_status") not in ("paid", "free"):
                        with db() as conn:
                            ev = conn.execute("SELECT * FROM events WHERE id = ?", (r["event_id"],)).fetchone()
                            u = conn.execute("SELECT role FROM users WHERE id = ?", (r["user_id"],)).fetchone()
                        u = row_to_dict(u) if u else {}
                        if ev:
                            order_ref = reg.get("wayforpay_order_reference") or f"ev{r['event_id']}_u{r['user_id']}_{int(datetime.now().timestamp())}"
                            role_u = (u.get("role") or "").strip()
                            base_price, discount_price = _event_prices_for_role(row_to_dict(ev), role_u)
                            amount = discount_price if discount_price is not None else base_price
                            if not reg.get("wayforpay_order_reference"):
                                with db() as conn:
                                    conn.execute("UPDATE registrations SET wayforpay_order_reference = ?, amount_uah = ?, updated_at = ? WHERE user_id = ? AND event_id = ?", (order_ref, amount, utcnow_iso(), r["user_id"], r["event_id"]))
                            invoice_url = await create_invoice(
                                order_ref, float(amount), f"Квиток: {ev['title']}",
                                client_email=r.get("email") or "",
                                client_phone=r.get("phone") or "",
                                client_name=r.get("full_name") or "",
                            )
                            evd = row_to_dict(ev)
                            role_u = (u.get("role") or "").strip()
                            if role_u == "member":
                                tpl = evd.get("wayforpay_member_url")
                            elif role_u == "lead":
                                tpl = evd.get("wayforpay_lead_url")
                            else:
                                tpl = evd.get("wayforpay_guest_url")
                            pay_url = _event_payment_url_from_template(
                                tpl or "",
                                order_ref=order_ref,
                                amount=int(amount or 0),
                                event_id=int(r["event_id"]),
                                user_id=int(r["user_id"]),
                            ) or invoice_url
                            if pay_url:
                                with db() as conn:
                                    conn.execute(
                                        "UPDATE registrations SET wayforpay_order_reference = ?, updated_at = ? WHERE user_id = ? AND event_id = ?",
                                        (order_ref, utcnow_iso(), r["user_id"], r["event_id"]),
                                    )
                                await send_message(r["tg_user_id"], (r.get("message_text") or "Нагадування") + "\n\nНове посилання для оплати:", reply_markup=inline_url_button("💳 Оплатити внесок", pay_url))
                            else:
                                await send_message(r["tg_user_id"], (r.get("message_text") or "Нагадування"))
                        else:
                            await send_message(r["tg_user_id"], (r.get("message_text") or "Нагадування"))
                    else:
                        await send_message(r["tg_user_id"], (r.get("message_text") or "Нагадування"))
                else:
                    await send_message(r["tg_user_id"], (r.get("message_text") or "Нагадування"))
                mark_sent(r["id"])
        except Exception:
            pass
        await asyncio.sleep(30)
